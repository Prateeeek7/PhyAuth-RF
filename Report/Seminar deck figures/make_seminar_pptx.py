"""Builds 'RF Fingerprinting Seminar.pptx': 16:9, light background, plain text + images,
boxes only for callouts. Also writes preview PNGs (preview/slide_NN.png) drawn from the
same layout so the result can be checked without PowerPoint."""
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Emu, Pt

HERE = Path(__file__).resolve().parent
REP = HERE.parent
FIG = REP / "Project report" / "figures"
AS = REP / "Project report" / "src" / "assets"
OUT = REP / "RF Fingerprinting Seminar.pptx"
PREV = HERE / "preview"
PREV.mkdir(exist_ok=True)
SF = HERE

C = {"bg": "F8F7F3", "alt": "EEF1F6", "dark": "1B2638", "body": "3F4A5C", "muted": "7A8394",
     "teal": "0E6E73", "orange": "B4531A", "tint_t": "E3EFEF", "tint_o": "F6E7DC", "line": "D9D4C7", "white": "FFFFFF"}
HEAD, TEXT = "Georgia", "Calibri"
PX = 12192000 / 1920          # EMU per design pixel; 1 px = 0.5 pt
FONTS = {"Georgia": "/System/Library/Fonts/Supplemental/Georgia.ttf",
         "Georgia-B": "/System/Library/Fonts/Supplemental/Georgia Bold.ttf",
         "Calibri": "/System/Library/Fonts/Supplemental/Arial.ttf",          # Arial is wider than Calibri: conservative preview
         "Calibri-B": "/System/Library/Fonts/Supplemental/Arial Bold.ttf"}

prs = Presentation()
prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)


def rgb(h):
    return RGBColor.from_string(h)


def px(v):
    return Emu(int(v * PX))


class Slide:
    def __init__(self, bg="bg"):
        self.s = prs.slides.add_slide(prs.slide_layouts[6])
        f = self.s.background.fill
        f.solid()
        f.fore_color.rgb = rgb(C[bg])
        self.img = Image.new("RGB", (1920, 1080), "#" + C[bg])
        self.d = ImageDraw.Draw(self.img)
        self.n = len(prs.slides)

    def rect(self, x, y, w, h, fill, radius=14):
        shp = self.s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                                      px(x), px(y), px(w), px(h))
        if radius:
            shp.adjustments[0] = min(0.5, radius / min(w, h))
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(C[fill])
        shp.line.fill.background()
        shp.shadow.inherit = False
        self.d.rounded_rectangle([x, y, x + w, y + h], radius=radius, fill="#" + C[fill])

    def picture(self, path, x, y, w, h):
        im = Image.open(path)
        sc = min(w / im.width, h / im.height)
        dw, dh = im.width * sc, im.height * sc
        ox, oy = x + (w - dw) / 2, y + (h - dh) / 2
        self.s.shapes.add_picture(str(path), px(ox), px(oy), px(dw), px(dh))
        self.img.paste(im.convert("RGB").resize((int(dw), int(dh))), (int(ox), int(oy)))

    def text(self, x, y, w, paras, size=18, color="body", font=TEXT, bold=False, align="l", spacing=1.18, gap=8):
        tb = self.s.shapes.add_textbox(px(x), px(y), px(w), px(40))
        tf = tb.text_frame
        tf.word_wrap = True
        for side in ("left", "right", "top", "bottom"):
            setattr(tf, f"margin_{side}", Emu(0))
        cy = y
        for i, p in enumerate(paras):
            opts = {}
            if isinstance(p, tuple):
                p, opts = p
            psize, pcol = opts.get("size", size), opts.get("color", color)
            pfont, pbold, pal = opts.get("font", font), opts.get("bold", bold), opts.get("align", align)
            pgap = opts.get("gap", gap)
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[pal]
            para.line_spacing = spacing
            para.space_after = Pt(pgap / 2)
            for k, seg in enumerate(p.split("**")):
                if seg:
                    r = para.add_run()
                    r.text = seg
                    r.font.size = Pt(psize)
                    r.font.name = pfont
                    r.font.bold = pbold or k % 2 == 1
                    r.font.color.rgb = rgb(C[pcol])
            f = ImageFont.truetype(FONTS[pfont + ("-B" if pbold else "")], int(psize * 2))
            line, lines = "", []
            for wd in p.replace("**", "").split(" "):
                t = (line + " " + wd).strip()
                if f.getlength(t) <= w or not line:
                    line = t
                else:
                    lines.append(line)
                    line = wd
            lines.append(line)
            for ln in lines:
                lw = f.getlength(ln)
                lx = x if pal == "l" else (x + (w - lw) / 2 if pal == "c" else x + w - lw)
                self.d.text((lx, cy), ln, font=f, fill="#" + C[pcol])
                cy += psize * 2 * spacing
            cy += pgap
        tb.height = px(max(cy - y, 20))
        return cy

    def header(self, eyebrow, title, eye="teal"):
        self.text(128, 70, 1664, [eyebrow.upper()], size=13, color=eye, bold=True)
        b = self.text(128, 104, 1664, [title], size=30, color="dark", font=HEAD, bold=True, spacing=1.05)
        self.rect(128, b + 10, 120, 6, eye, radius=0)
        return b + 50

    def footer(self, note="RF Fingerprinting · Seminar"):
        self.text(128, 1010, 1300, [note], size=11, color="muted")
        self.text(1672, 1010, 120, [str(self.n)], size=11, color="muted", align="r")

    def notes(self, t):
        self.s.notes_slide.notes_text_frame.text = t

    def done(self, note):
        self.notes(note)
        self.img.save(PREV / f"slide_{self.n:02d}.png")


def bullets(items, size=19):
    return [("•  " + t, {"size": size}) for t in items]


def paper(sl, x, y, w, cite, title, items):
    return sl.text(x, y, w, [(cite, {"size": 14, "color": "orange", "bold": True, "gap": 4}),
                             (title, {"size": 21, "color": "dark", "bold": True, "gap": 10})] + bullets(items)) + 26


def callout(sl, x, y, w, txt, tint="tint_t", bar="teal", size=19):
    f = ImageFont.truetype(FONTS["Calibri"], size * 2)
    lines = int(f.getlength(txt.replace("**", "")) // (w - 64)) + 1
    h = lines * size * 2 * 1.18 + 36
    sl.rect(x, y, w, h, tint, radius=10)
    sl.rect(x, y, 8, h, bar, radius=0)
    sl.text(x + 32, y + 18, w - 64, [txt], size=size, color="dark")
    return y + h


# 1 cover -----------------------------------------------------------------------
s = Slide()
s.rect(0, 0, 1920, 16, "teal", radius=0)
s.picture(AS / "vit_logo.png", 1400, 70, 392, 130)
y = s.text(128, 260, 1600, [("SEMINAR  ·  BECE306P DIGITAL COMMUNICATION SYSTEMS", {"size": 15, "color": "teal", "bold": True, "gap": 24}),
                            ("RF Fingerprinting for Physical-Layer Device Authentication", {"size": 46, "color": "dark", "font": HEAD, "bold": True, "gap": 28}),
                            ("A survey of how radios are identified by their hardware imperfections, and a low-cost implementation on a 433 MHz ESP32 receiver",
                             {"size": 22})], spacing=1.08)
s.rect(128, y + 40, 1664, 2, "line", radius=0)
s.text(128, y + 80, 900, [("PRESENTED BY", {"size": 13, "color": "orange", "bold": True}),
                          "Pratik Kumar (24BML0027)  ·  Nikunj Goydani (24BEC0338)",
                          "Shruti (24BML0033)  ·  Uday Krishan Khandelwal (24BML0017)"], size=18)
s.text(1150, y + 80, 640, [("GUIDE", {"size": 13, "color": "orange", "bold": True}), "Dr. NandaKumar", "SENSE, VIT"], size=18)
s.done("This seminar surveys RF fingerprinting: identifying a radio from imperfections in its hardware rather than from its data. "
       "We cover the main families of published methods, compare them, find the gap, and in the last four slides show our implementation.")

# 2 motivation ------------------------------------------------------------------
s = Slide()
y0 = s.header("Motivation", "Authenticating the message is not authenticating the device")
s.text(128, y0, 800, [("The problem", {"size": 22, "color": "dark", "bold": True, "gap": 14})] + bullets([
    "Security usually checks a code or key carried in the data",
    "Cheap IoT links (433 MHz remotes, sensors) send fixed or weakly protected codes",
    "Anyone can record the signal and replay it with a $5 transmitter",
    "The receiver cannot tell the genuine radio from a copy"]))
s.text(1040, y0, 752, [("Known attacks", {"size": 22, "color": "dark", "bold": True, "gap": 14}),
                       ("**RollJam** (Kamkar, DEF CON 2015): jams and records a rolling code, replays it later.", {"gap": 18}),
                       ("**Garcia et al.** (USENIX Security 2016): keyless entry of many vehicles broken by eavesdropping and cryptanalysis.", {})], size=19)
callout(s, 128, 820, 1664, "**Idea:** verify who transmitted from the radio hardware itself, not from the code it sends.", "tint_o", "orange", 21)
s.footer()
s.done("Most wireless security authenticates the data. Fixed codes can be recorded and replayed; RollJam and Garcia et al. show even rolling "
       "codes and car keys are vulnerable. Physical-layer authentication asks whether the signal came from the expected transmitter.")

# 3 fingerprint -------------------------------------------------------------------
s = Slide()
y0 = s.header("Background", "What is an RF fingerprint?")
s.text(128, y0, 700, [("Small hardware imperfections, unique to each transmitter", {"size": 21, "color": "dark", "bold": True, "gap": 16})] + bullets([
    "Oscillator frequency offset and clock skew", "Phase noise and timing jitter", "Amplifier non-linearity, turn-on transient",
    "I/Q imbalance and DC offset", "Filter and DAC distortion"]) +
       [("From component tolerances: stable for one device, different across devices of the same model.", {"size": 17, "color": "muted"})])
s.picture(SF / "fig_impairments.png", 880, y0 - 10, 912, 540)
callout(s, 128, 830, 810, "**Transient-based:** analyse the turn-on of a burst")
callout(s, 982, 830, 810, "**Steady-state:** analyse errors in the modulated part")
s.footer("RF Fingerprinting · Seminar · taxonomy after Danev, Zanetti & Čapkun, ACM Computing Surveys 2012")
s.done("Imperfections enter at every stage of the transmitter. Methods split into transient-based and steady-state families.")


def survey(eyebrow, title, papers, take, img, note):
    s = Slide()
    y = s.header(eyebrow, title)
    top = y
    for cite, t, items in papers:
        y = paper(s, 128, y, 760, cite, t, items)
    callout(s, 128, max(y, 790), 760, take)
    s.picture(SF / img, 940, top, 852, 920 - top)
    s.footer()
    s.done(note)


survey("Related work · 1", "Clock skew: every crystal ticks at its own rate", [
    ("Kohno, Broido & Claffy · IEEE TDSC 2005", "Remote physical device fingerprinting", ["Clock skew from TCP timestamps", "Stable per device, measured remotely"]),
    ("Jana & Kasera · IEEE TMC 2010", "Rogue access-point detection", ["Skew from Wi-Fi beacon timestamps", "Catches an AP copying a real one"])],
    "**Takeaway:** timing alone carries identity, but needs long observation.", "fig_clockskew.png",
    "Every crystal runs slightly fast or slow; the offset grows linearly and its slope identifies the device. Kohno used TCP timestamps; "
    "Jana and Kasera used Wi-Fi beacons to detect rogue access points.")
survey("Related work · 2", "Turn-on transients", [
    ("Ureten & Serinken · Canadian J. ECE 2007", "Wireless security through RF fingerprinting", ["Find where the transient starts", "Classify its shape with a neural network"]),
    ("Danev & Čapkun · ACM/IEEE IPSN 2009", "Identifying identical sensor nodes", ["IEEE 802.15.4 nodes told apart", "High-rate oscilloscope in the lab"])],
    "**Limitation:** microsecond transients need very fast sampling.", "fig_transient.png",
    "Transient methods look at the turn-on period before data starts. They separate even identical nodes but need very fast sampling.")
survey("Related work · 3", "Modulation errors: PARADIS", [
    ("Brik, Banerjee, Gruteser & Oh · ACM MobiCom 2008", "Radiometric signatures of 802.11b cards",
     ["Frequency error, SYNC correlation, I/Q offset", "Magnitude and phase error", "k-NN and SVM classifiers", "Vector signal analyser for capture"])],
    "**Strength:** whole frame.  **Limit:** lab-grade equipment.", "fig_constellation.png",
    "PARADIS measures modulation errors of 802.11b cards (offset, rotation, I/Q imbalance in the constellation) and classifies them with "
    "k-NN and SVM. Robust, but it needs a vector signal analyser.")

# 7 deep learning ------------------------------------------------------------------
s = Slide()
y0 = s.header("Related work · 4", "Deep learning on raw I/Q samples")
for i, (cite, t) in enumerate([("Merchant et al. · IEEE JSTSP 2018", "CNN on the error signal of ZigBee devices"),
                               ("Riyaz et al. · IEEE Comm. Mag. 2018", "CNN learns features from raw I/Q of SDR radios"),
                               ("Sankhe et al. · INFOCOM 2019 (ORACLE)", "Same-model USRP radios; injected impairments help")]):
    s.text(128 + i * 565, y0, 520, [(cite, {"size": 14, "color": "orange", "bold": True, "gap": 6}), (t, {"size": 19})])
s.picture(SF / "fig_cnn.png", 128, y0 + 150, 1664, 320)
callout(s, 128, 850, 810, "**Gain:** no hand-made features; high accuracy")
callout(s, 982, 850, 810, "**Cost:** SDR, large models, GPU; channel-sensitive", "tint_o", "orange")
s.footer()
s.done("From 2018 CNNs learned fingerprints directly from I/Q samples (Merchant, Riyaz, ORACLE). High accuracy, but they need SDRs and GPUs "
       "and are sensitive to the channel.")

survey("Related work · 5", "The channel problem", [
    ("Shen, Zhang, Marshall & Cavallaro · IEEE TIFS 2022", "Channel-robust identification for LoRa",
     ["Multipath can change the signal more than the device", "Channel-independent spectrogram", "Training with simulated channels"])],
    "**Lesson:** features must separate the device from the channel and the receiver.", "fig_channel.png",
    "Multipath can dominate the device response; Shen et al. removed most of it with a channel-independent spectrogram. "
    "In our project the cheap receiver's own distortion plays the same role.")
survey("Related work · 6", "Rejecting the unknown attacker", [
    ("Schroff et al. CVPR 2015 · Hermans et al. 2017", "Metric learning", ["Triplet loss: same device close, others far"]),
    ("Lee et al. NeurIPS 2018 · Ledoit & Wolf 2004", "Mahalanobis distance", ["Distance to each enrolled device", "Reject beyond a threshold"])],
    "**Why:** softmax always picks a known device, even for an attacker.", "fig_openset.png",
    "Closed-set classifiers label an attacker as some enrolled device; metric learning plus a Mahalanobis threshold per device rejects unknowns.")

# 10 comparison --------------------------------------------------------------------
s = Slide("alt")
y0 = s.header("Comparison", "The surveyed work side by side")
rows = [("Work", "Capture", "Feature and method", "Limitation"),
        ("Kohno 2005; Jana 2010", "Network timestamps", "Clock skew", "Long observation"),
        ("Ureten 2007; Danev 2009", "High-rate oscilloscope", "Turn-on transient", "Very fast sampling"),
        ("Brik 2008 (PARADIS)", "Signal analyser", "Modulation errors, k-NN / SVM", "Lab equipment"),
        ("Merchant, Riyaz 2018", "SDR I/Q", "CNN", "SDR and GPU"),
        ("Sankhe 2019 (ORACLE)", "USRP SDR", "CNN, injected impairments", "Costly testbed"),
        ("Shen 2022", "SDR, LoRa", "Channel-independent spectrogram", "SDR required"),
        ("Lee 2018 (ML)", "Deep features", "Mahalanobis open-set", "Not RF-specific")]
widths, rh = (430, 380, 540, 314), 76
tbl = s.s.shapes.add_table(len(rows), 4, px(128), px(y0), px(1664), px(rh * len(rows))).table
for j, w in enumerate(widths):
    tbl.columns[j].width = px(w)
for i, r in enumerate(rows):
    tbl.rows[i].height = px(rh)
    fill = C["dark"] if i == 0 else (C["white"] if i % 2 else C["bg"])
    s.d.rectangle([128, y0 + i * rh, 1792, y0 + (i + 1) * rh], fill="#" + fill)
    x = 128
    for j, v in enumerate(r):
        cell = tbl.cell(i, j)
        cell.text = v
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        run = cell.text_frame.paragraphs[0].runs[0]
        run.font.size, run.font.name, run.font.bold = Pt(16), TEXT, i == 0
        run.font.color.rgb = rgb(C["bg"] if i == 0 else C["dark"])
        cell.fill.solid()
        cell.fill.fore_color.rgb = rgb(fill)
        f = ImageFont.truetype(FONTS["Calibri-B" if i == 0 else "Calibri"], 32)
        s.d.text((x + 16, y0 + i * rh + 20), v, font=f, fill="#" + (C["bg"] if i == 0 else C["dark"]))
        x += widths[j]
s.footer()
s.done("Nearly every RF result depends on expensive capture hardware; the open-set methods are general, not RF-specific.")

# 11 gap ---------------------------------------------------------------------------
s = Slide()
y0 = s.header("Research gap", "What the literature leaves open")
for i, (t, d) in enumerate([("Cost", "Capture hardware costs far more than the device it protects."),
                            ("Cheap receivers distort", "Low-cost OOK receivers bias pulse widths; methods assume clean I/Q."),
                            ("Not on the device", "Decisions run on a PC or GPU, not on the receiver itself."),
                            ("Single factor", "The fingerprint is rarely combined with a user secret.")]):
    x, y = 128 + (i % 2) * 860, y0 + (i // 2) * 230
    s.text(x, y, 790, [(f"0{i + 1}", {"size": 28, "color": "teal", "font": HEAD, "bold": True, "gap": 6}),
                       (t, {"size": 22, "color": "dark", "bold": True, "gap": 6}), (d, {"size": 19})])
callout(s, 128, 820, 1664, "**Question:** can a $5 microcontroller behind a $1 receiver authenticate the transmitter on its own?", "tint_o", "orange", 22)
s.footer()
s.done("Four gaps: cost, receiver distortion, no on-device decision, single factor. That leads to our question.")

# 12 our architecture -----------------------------------------------------------------
s = Slide()
y0 = s.header("Our work · PhyAuth-RF", "A 433 MHz fingerprint verifier on an ESP32", "orange")
s.picture(FIG / "fig_architecture.png", 128, y0, 1100, 940 - y0)
s.picture(FIG / "fig_receiver_labelled.png", 1260, y0, 532, 320)
s.text(1260, y0 + 350, 532, bullets(["Genuine: QC5681 EV1527 remote", "Receiver: ESP32 + SRX882 (≈ Rs. 650)",
                                     "Attacker: ESP32-S3 + STX882", "Every decision on the ESP32"], 18))
s.footer()
s.done("PhyAuth-RF: a QC5681 remote, an ESP32 + SRX882 receiver, and an ESP32-S3 + STX882 emulator and attacker. "
       "All decisions are made on the ESP32; the laptop only trains and monitors.")

# 13 our method -----------------------------------------------------------------------
s = Slide()
y0 = s.header("Our work · method", "Timing fingerprints through a distorting receiver", "orange")
s.picture(FIG / "fig_pipeline.png", 128, y0, 900, 940 - y0)
y = y0 + 10
for t, d in [("4.17 ns edge timing", "CPU cycle counter, no SDR or ADC"),
             ("Distortion-proof clock", "Pair period 4T + 31T sync cancel receiver bias"),
             ("11 features → EdgeMLP", "963 parameters, 7 KB on the chip"),
             ("Open-set decision", "Mahalanobis score; median of 4 frames ≤ 1")]:
    y = s.text(1090, y, 702, [(t, {"size": 22, "color": "dark", "bold": True, "gap": 4}), (d, {"size": 19, "gap": 0})]) + 40
s.footer()
s.done("Edge timestamps at 4.17 ns; clock from the pair period and sync gap; 11 features into a 963-parameter EdgeMLP; "
       "Mahalanobis decision over four frames.")

# 14 our results ----------------------------------------------------------------------
s = Slide("alt")
y0 = s.header("Our work · results", "Clones rejected, and a two-factor smart lock", "orange")
stats = [("100 %", "closed-set accuracy", "teal"), ("0.49 %", "frame EER", "teal"), ("30 / 30", "clone bursts rejected", "orange"),
         ("30 / 30", "unknown bursts rejected", "orange"), ("≈ 172 ms", "decision time", "teal"), ("7 KB", "model on the ESP32", "teal")]
for i, (v, l, c) in enumerate(stats):
    x, y = 128 + (i % 2) * 400, y0 + (i // 2) * 215
    s.rect(x, y, 380, 190, "bg", radius=14)
    s.text(x + 28, y + 28, 330, [(v, {"size": 31, "color": c, "font": HEAD, "bold": True, "gap": 8}), (l, {"size": 16})])
s.picture(FIG / "fig_lock_fsm.png", 960, y0, 832, 390)
s.text(960, y0 + 420, 832, [("Two-factor smart lock", {"size": 22, "color": "dark", "bold": True, "gap": 8}),
                            ("Opens only for the **right remote** (RF fingerprint) and the **right ON/OFF pattern**. "
                             "A clone with the correct pattern still fails; 3 failures cause a lockout.", {"size": 19})])
s.footer("RF Fingerprinting · Seminar · held-out real captures")
s.done("100% accuracy, 0.49% frame EER, all clone and unknown bursts rejected, about 172 ms, 7 KB. Two-factor smart lock on the same ESP32. "
       "Limits: one genuine remote, small test sets, 0.75 to 1.25 m range.")

# 15 conclusion -------------------------------------------------------------------------
s = Slide()
y0 = s.header("Conclusion", "From lab instruments to a $5 microcontroller")
s.text(128, y0, 800, [("Key points", {"size": 22, "color": "dark", "bold": True, "gap": 14})] + bullets([
    "Hardware imperfections identify a transmitter beyond its data",
    "Research moved from clock skew and transients to deep learning on SDR I/Q",
    "Open problems: cost, channel and receiver effects, unknown attackers",
    "PhyAuth-RF runs a timing fingerprint on-chip at low cost"]))
s.text(1000, y0, 792, [("References", {"size": 22, "color": "dark", "bold": True, "gap": 12})] + [(r, {"size": 15, "gap": 2}) for r in [
    "Kohno et al., IEEE TDSC 2005", "Jana & Kasera, IEEE TMC 2010", "Ureten & Serinken, Can. J. ECE 2007",
    "Danev & Čapkun, IPSN 2009", "Brik et al., ACM MobiCom 2008", "Danev et al., ACM Comput. Surv. 2012",
    "Merchant et al., IEEE JSTSP 2018", "Riyaz et al., IEEE Commun. Mag. 2018", "Sankhe et al., IEEE INFOCOM 2019",
    "Shen et al., IEEE TIFS 2022", "Schroff et al., CVPR 2015;  Lee et al., NeurIPS 2018", "Garcia et al., USENIX Security 2016"]])
s.text(128, 870, 1000, ["Thank you. Questions?"], size=32, color="dark", font=HEAD, bold=True)
s.footer()
s.done("Summary and thanks.")

prs.save(OUT)
print("wrote", OUT, len(prs.slides), "slides")
