"""Small python-docx toolkit: numbered headings, figures, tables, equations, TOC."""
import re
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
plt.rcParams["font.family"] = "Times New Roman"
plt.rcParams["mathtext.fontset"] = "custom"
plt.rcParams["mathtext.rm"] = "Times New Roman"
plt.rcParams["mathtext.it"] = "Times New Roman:italic"
plt.rcParams["mathtext.bf"] = "Times New Roman:bold"
plt.rcParams["mathtext.sf"] = "Times New Roman"
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

NAVY = RGBColor(0x0F, 0x29, 0x4A)
BLUE = RGBColor(0x1E, 0x56, 0xA0)
GREY = RGBColor(0x55, 0x5F, 0x6B)
FONT = "Times New Roman"
MONO = "Times New Roman"


def _set_cell_bg(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def _set_table_borders(table, color="7F8C8D", size=4):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    tblPr.append(borders)


def _fix_widths(table, widths_cm):
    table.autofit = False
    tblPr = table._tbl.tblPr
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tblPr.append(lay)
    tw = OxmlElement("w:tblW")
    tw.set(qn("w:w"), str(int(sum(widths_cm) * 567)))
    tw.set(qn("w:type"), "dxa")
    for old in tblPr.findall(qn("w:tblW")):
        tblPr.remove(old)
    tblPr.append(tw)
    grid = table._tbl.tblGrid
    for j, gc in enumerate(grid.findall(qn("w:gridCol"))):
        gc.set(qn("w:w"), str(int(widths_cm[j] * 567)))
    for row in table.rows:
        for j, w in enumerate(widths_cm):
            row.cells[j].width = Cm(w)


def _no_borders(table):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tblPr.append(borders)


def _field(run, instr):
    for kind, text in (("begin", None), (None, instr), ("separate", None), (None, "1"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
            run._r.append(el)
        elif text == instr:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = instr
            run._r.append(el)
        else:
            el = OxmlElement("w:t")
            el.text = text
            run._r.append(el)


class Report:
    def __init__(self, template: Path, workdir: Path):
        self.doc = Document(str(template))
        self.work = workdir
        self.work.mkdir(exist_ok=True)
        body = self.doc.element.body
        for child in list(body):
            if child.tag != qn("w:sectPr"):
                body.remove(child)
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        sec.left_margin = sec.right_margin = Cm(2.3)
        sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.0)
        sec.header_distance = sec.footer_distance = Cm(1.0)
        st = self.doc.styles["Normal"]
        st.font.name = FONT
        st.font.size = Pt(11.5)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        pf = st.paragraph_format
        pf.space_after = Pt(6)
        pf.line_spacing = 1.15
        for name, size, color in (("Heading 1", 17, NAVY), ("Heading 2", 14, BLUE), ("Heading 3", 12.5, NAVY)):
            h = self.doc.styles[name]
            h.font.name = FONT
            h.font.size = Pt(size)
            h.font.bold = True
            h.font.italic = False
            h.font.color.rgb = color
            rpr = h.element.get_or_add_rPr()
            rf = rpr.find(qn("w:rFonts"))
            if rf is None:
                rf = OxmlElement("w:rFonts")
                rpr.append(rf)
            for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
                rf.set(qn(a), FONT)
            for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                if rf.get(qn(a)) is not None:
                    del rf.attrib[qn(a)]
            h.paragraph_format.space_before = Pt(14 if name != "Heading 1" else 0)
            h.paragraph_format.space_after = Pt(6 if name != "Heading 1" else 12)
            h.paragraph_format.keep_with_next = True
        self.chapter = 0
        self.sec = 0
        self.sub = 0
        self.fig_n = 0
        self.tab_n = 0
        self.eq_n = 0
        self.toc = []        # (level, text)
        self.figs = []       # caption text
        self.tabs = []
        self.eq_cache = 0

    # ------------------------------------------------------------------ text
    def para(self, text="", align="justify", size=None, bold=False, italic=False, color=None,
             space_after=6, space_before=0, keep=False, font=None, indent=None, line=None):
        p = self.doc.add_paragraph()
        p.alignment = {"justify": WD_ALIGN_PARAGRAPH.JUSTIFY, "center": WD_ALIGN_PARAGRAPH.CENTER,
                       "left": WD_ALIGN_PARAGRAPH.LEFT, "right": WD_ALIGN_PARAGRAPH.RIGHT}[align]
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.space_before = Pt(space_before)
        if line:
            p.paragraph_format.line_spacing = line
        if keep:
            p.paragraph_format.keep_with_next = True
        if indent is not None:
            p.paragraph_format.left_indent = Cm(indent)
        self._runs(p, text, size=size, bold=bold, italic=italic, color=color, font=font)
        return p

    def _runs(self, p, text, size=None, bold=False, italic=False, color=None, font=None):
        # **bold**, *italic*, `code`
        for tok in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`|_\{[^}]+\}|\^\{[^}]+\})", text):
            if not tok:
                continue
            r = p.add_run()
            if tok.startswith("**"):
                r.text, r.bold = tok[2:-2], True
            elif tok.startswith("*"):
                r.text, r.italic = tok[1:-1], True
            elif tok.startswith("_{") or tok.startswith("^{"):
                r.text = tok[2:-1]
                r.font.subscript = tok[0] == "_"
                r.font.superscript = tok[0] == "^"
            elif tok.startswith("`"):
                r.text = tok[1:-1]
                r.font.name = MONO
                r.font.size = Pt((size or 12) - 1.5)
            else:
                r.text = tok
            if bold:
                r.bold = True
            if italic:
                r.italic = True
            if size and not tok.startswith("`"):
                r.font.size = Pt(size)
            if color is not None:
                r.font.color.rgb = color
            if font and not tok.startswith("`"):
                r.font.name = font

    def bullets(self, items, numbered=False, size=None):
        for i, it in enumerate(items, 1):
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.left_indent = Cm(0.9)
            p.paragraph_format.first_line_indent = Cm(-0.5)
            p.paragraph_format.space_after = Pt(3)
            lead = f"{i}." if numbered else "•"
            r = p.add_run(lead + "\t")
            p.paragraph_format.tab_stops.add_tab_stop(Cm(0.9))
            self._runs(p, it, size=size)

    def page_break(self):
        p = self.doc.add_paragraph()
        p.add_run().add_break(WD_BREAK.PAGE)
        p.paragraph_format.space_after = Pt(0)

    # -------------------------------------------------------------- headings
    def h1(self, title, numbered=True, toc=True, new_page=True):
        first = None
        if numbered:
            self.chapter += 1
            self.sec = self.fig_n = self.tab_n = self.eq_n = 0
            label = f"Chapter {self.chapter}"
            p = self.doc.add_paragraph(style="Heading 1")
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(label.upper())
            r.font.size = Pt(12)
            r.font.color.rgb = GREY
            text = f"{label}: {title}"
            first = p
        else:
            text = title
        p = self.doc.add_paragraph(style="Heading 1")
        p.add_run(title)
        if new_page:
            (first or p).paragraph_format.page_break_before = True
        self._rule(p)
        if toc:
            self.toc.append((1, text if numbered else title))

    def _rule(self, p):
        pPr = p._p.get_or_add_pPr()
        bdr = OxmlElement("w:pBdr")
        b = OxmlElement("w:bottom")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), "8")
        b.set(qn("w:space"), "4")
        b.set(qn("w:color"), "1E56A0")
        bdr.append(b)
        pPr.append(bdr)

    def h2(self, title):
        self.sec += 1
        self.sub = 0
        t = f"{self.chapter}.{self.sec}  {title}"
        self.doc.add_paragraph(t, style="Heading 2")
        self.toc.append((2, t))

    def h3(self, title):
        self.sub += 1
        t = f"{self.chapter}.{self.sec}.{self.sub}  {title}"
        self.doc.add_paragraph(t, style="Heading 3")

    # --------------------------------------------------------------- objects
    def figure(self, path, caption, width_cm=16.0):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        p.add_run().add_picture(str(path), width=Cm(width_cm))
        self.fig_n += 1
        label = f"Figure {self.chapter}.{self.fig_n}"
        c = self.doc.add_paragraph()
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.paragraph_format.space_after = Pt(10)
        r = c.add_run(label + ": ")
        r.bold = True
        r.font.size = Pt(10.5)
        self._runs(c, caption, size=10.5, italic=False)
        self.figs.append(f"{label}: {caption}".replace("*", "").replace("`", ""))
        return label

    def next_fig(self, k=1):
        return f"Figure {self.chapter}.{self.fig_n + k}"

    def next_tab(self, k=1):
        return f"Table {self.chapter}.{self.tab_n + k}"

    def table(self, header, rows, caption, widths=None, size=10, align=None, bold_first=False, numbered=True):
        if numbered:
            self.tab_n += 1
        label = f"Table {self.chapter}.{self.tab_n}" if numbered else ""
        c = self.doc.add_paragraph()
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.paragraph_format.keep_with_next = True
        c.paragraph_format.space_before = Pt(6)
        c.paragraph_format.space_after = Pt(4)
        if numbered:
            r = c.add_run(label + ": ")
            r.bold = True
            r.font.size = Pt(10.5)
            self.tabs.append(f"{label}: {caption}".replace("*", "").replace("`", ""))
        self._runs(c, caption, size=10.5)
        t = self.doc.add_table(rows=1 + len(rows), cols=len(header))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        _set_table_borders(t)
        for j, h in enumerate(header):
            cell = t.rows[0].cells[j]
            cell.text = ""
            pp = cell.paragraphs[0]
            pp.paragraph_format.space_after = Pt(1)
            pp.paragraph_format.line_spacing = 1.0
            rr = pp.add_run(h)
            rr.bold = True
            rr.font.size = Pt(size)
            rr.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            _set_cell_bg(cell, "0F294A")
        for i, row in enumerate(rows, 1):
            for j, val in enumerate(row):
                cell = t.rows[i].cells[j]
                cell.text = ""
                pp = cell.paragraphs[0]
                pp.paragraph_format.space_after = Pt(1)
                pp.paragraph_format.line_spacing = 1.0
                if align and align[j] == "c":
                    pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                elif align and align[j] == "r":
                    pp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                self._runs(pp, str(val), size=size, bold=(bold_first and j == 0))
                if i % 2 == 0:
                    _set_cell_bg(cell, "EEF3FA")
        # keep rows together where possible
        for row in t.rows:
            trPr = row._tr.get_or_add_trPr()
            cs = OxmlElement("w:cantSplit")
            trPr.append(cs)
            for cell in row.cells:
                for pp in cell.paragraphs:
                    pp.paragraph_format.keep_with_next = True
        for cell in t.rows[-1].cells:
            for pp in cell.paragraphs:
                pp.paragraph_format.keep_with_next = False
        if widths:
            _fix_widths(t, widths)
        sp = self.doc.add_paragraph()
        sp.paragraph_format.space_after = Pt(4)
        sp.paragraph_format.line_spacing = 0.8
        return label

    def equation(self, latex, fontsize=13):
        self.eq_n += 1
        self.eq_cache += 1
        png = self.work / f"eq_{self.eq_cache:03d}.png"
        fig = plt.figure(figsize=(0.01, 0.01))
        fig.text(0, 0, f"${latex}$", fontsize=fontsize)
        fig.savefig(png, dpi=300, bbox_inches="tight", pad_inches=0.04, transparent=False, facecolor="white")
        plt.close(fig)
        from PIL import Image
        w_px, h_px = Image.open(png).size
        t = self.doc.add_table(rows=1, cols=2)
        _no_borders(t)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        _fix_widths(t, [15.0, 1.4])
        c0, c1 = t.rows[0].cells
        p0 = c0.paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p0.paragraph_format.space_after = Pt(2)
        w_cm, h_cm = w_px / 300 * 2.54, h_px / 300 * 2.54
        k = min(1.0, 2.4 / h_cm, 14.8 / w_cm)
        p0.add_run().add_picture(str(png), width=Cm(w_cm * k))
        p1 = c1.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = p1.add_run(f"({self.chapter}.{self.eq_n})")
        r.font.size = Pt(11)
        from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
        c0.vertical_alignment = c1.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        return f"({self.chapter}.{self.eq_n})"

    def code(self, text, size=9):
        t = self.doc.add_table(rows=1, cols=1)
        _set_table_borders(t, "AAB7B8", 4)
        _fix_widths(t, [16.4])
        cell = t.rows[0].cells[0]
        _set_cell_bg(cell, "F4F6F7")
        cell.text = ""
        p = cell.paragraphs[0]
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(0)
        for k, line in enumerate(text.strip("\n").split("\n")):
            r = p.add_run(line)
            r.font.name = MONO
            r.font.size = Pt(size)
            if k < len(text.strip("\n").split("\n")) - 1:
                r.add_break()
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # ------------------------------------------------------- header / footer
    def header_footer(self, left, right):
        sec = self.doc.sections[0]
        sec.different_first_page_header_footer = True
        hp = sec.header.paragraphs[0]
        hp.text = ""
        hp.paragraph_format.tab_stops.add_tab_stop(Cm(16.4), WD_TAB_ALIGNMENT.RIGHT)
        r = hp.add_run(f"{left}\t{right}")
        r.font.size = Pt(9)
        r.font.color.rgb = GREY
        self._rule(hp)
        fp = sec.footer.paragraphs[0]
        fp.text = ""
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = fp.add_run()
        r.font.size = Pt(10)
        _field(r, "PAGE")

    # ------------------------------------------------------------------ TOC
    def toc_block(self, entries, pages):
        """entries: list of (level, text); pages: dict text -> page (or None)."""
        for level, text in entries:
            p = self.doc.add_paragraph()
            p.paragraph_format.space_after = Pt(0 if level == 2 else (1 if level == 3 else 1))
            p.paragraph_format.space_before = Pt(2 if level == 1 else 0)
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.left_indent = Cm(0 if level == 1 else 0.8)
            p.paragraph_format.tab_stops.add_tab_stop(Cm(16.4), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
            r = p.add_run(text)
            r.font.size = Pt({1: 10.5, 2: 10, 3: 9.5}[level])
            r.bold = level == 1
            r2 = p.add_run("\t" + str(pages.get(text, "00")))
            r2.font.size = Pt(10 if level < 3 else 9.5)
            r2.bold = level == 1

    def save(self, path):
        self.doc.save(str(path))
