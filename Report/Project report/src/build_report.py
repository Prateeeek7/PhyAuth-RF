"""Builds 'PhyAuth-RF Project Report.docx' (and the PDF via Pages) from the
course template (chapter structure) and the seminar-style front page.

    python3 src/build_report.py            # three passes, page numbers resolved
"""
import json
import subprocess
import sys
from pathlib import Path

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

from docx_kit import Report, NAVY, BLUE, GREY

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
FIG = ROOT / "figures"
AS = HERE / "assets"
WORK = HERE / "_work"
OUT_DOCX = ROOT / "PhyAuth-RF Project Report.docx"
OUT_PDF = ROOT / "PhyAuth-RF Project Report.pdf"
REP = json.loads((HERE / "model_report.json").read_text())
ANA = json.loads((ROOT.parent / "Seminar" / "analysis.json").read_text())
XDAY = json.loads((ROOT.parent / "Seminar" / "crossday.json").read_text())

TITLE = ("PhyAuth-RF: Physical-Layer Device Authentication of 433 MHz Transmitters "
         "via RF Timing Fingerprints, Edge TinyML and a Two-Factor Smart Lock")
TEAM = [("Pratik Kumar", "24BML0027"), ("Nikunj Goydani", "24BEC0338"),
        ("Shruti", "24BML0033"), ("Uday Krishan Khandelwal", "24BML0017")]

E = REP["models"]["edge"]["verification"]
F = REP["models"]["fusion"]["verification"]


def pct(x, d=1):
    return f"{x * 100:.{d}f} %"


# ============================================================================
def build(prev, pages):
    """prev: lists (toc, figs, tabs) from the previous pass; pages: text -> page."""
    R = Report(AS / "template.docx", WORK)
    d = R.doc

    # ------------------------------------------------------------ cover page
    R.para("", space_after=30)
    R.para(TITLE, align="center", size=21, bold=True, color=RGBColor(0, 0, 0), line=1.15, space_after=34)
    R.para("PROJECT REPORT", align="center", size=16, bold=True, space_after=22)
    R.para("Bachelor of Technology", align="center", size=15, bold=True, space_after=4)
    R.para("School of Electronics Engineering (SENSE)", align="center", size=15, bold=True, space_after=4)
    R.para("BECE306P – Digital Communication Systems", align="center", size=12.5, bold=True, space_after=30)
    R.para("by", align="center", size=13, italic=True, space_after=10)
    for name, reg in TEAM:
        R.para(f"{name} ({reg})", align="center", size=13, bold=True, space_after=3)
    R.para("", space_after=18)
    R.para("Under the guidance of", align="center", size=13, bold=True, space_after=8)
    R.para("Dr. NandaKumar", align="center", size=15, bold=True, space_after=2)
    R.para("School of Electronics Engineering (SENSE)", align="center", size=11.5, space_after=46)
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(AS / "vit_logo.png"), width=Cm(7.5))
    R.para("October 2026", align="center", size=13, bold=True, space_before=14)

    R.chapter = 0

    # --------------------------------------------------------------- abstract
    R.h1("Abstract", numbered=False)
    R.para("Low-cost 433 MHz remote controls (garage doors, gates, alarm panels, smart plugs) transmit a fixed "
           "24-bit code with on–off keying. Anyone who records the code can replay it, because the receiver "
           "authenticates the *message*, never the *transmitter*. This project, **PhyAuth-RF**, authenticates the "
           "physical transmitter instead, using only the hardware a student can buy for about Rs. 2,000: two ESP32 "
           "boards, an SRX882 superheterodyne receiver, an STX882 transmitter and a QC5681 two-button EV1527 remote.")
    R.para("The receiver ESP32 timestamps every edge of the demodulated signal with its 240 MHz cycle counter "
           "(4.17 ns resolution). A distortion-invariant clock-recovery method (pair period and 31T sync gap) decodes the "
           "EV1527 frame even when the SRX882 slicer shortens every HIGH by up to a third of a symbol, and eleven "
           "code-independent timing features are extracted per frame. A 963-parameter EdgeMLP maps the features to a "
           "16-dimensional embedding, and an open-set Mahalanobis authenticator with Ledoit–Wolf covariance accepts "
           "or rejects each press on-chip, fusing four frames. A larger 21 k-parameter host model (1D ResNet + MLP) runs "
           "on the laptop in parallel, and a FastAPI + Next.js dashboard records data, trains, exports the model as a "
           "7 KB C header and visualises every decision.")
    R.para(f"On held-out real captures the deployed on-chip model reached 100 % closed-set accuracy, a frame-level "
           f"EER of {pct(E['frame']['eer'], 2)} (AUC {E['frame']['auc']:.5f}) and a burst-level EER of 0 %, with **all 30 "
           f"replay bursts of a code-cloning transmitter and all 30 bursts of an unknown transmitter rejected**, and no "
           f"genuine burst rejected. Fusing four frames cut the genuine rejection rate from 45.6 % (single frame) to 1.1 % "
           f"at a decision latency of about 172 ms. On top of the fingerprint, a two-factor **smart lock** runs on the "
           f"ESP32: the door opens only when the enrolled remote (something you *have*, verified by its RF fingerprint) "
           f"enters a secret ON/OFF press pattern (something you *know*), with lockout after three failures, learn mode, "
           f"flash persistence, a green status LED and a buzzer. The report documents the design, the measured limits of "
           f"the hardware (reliable window 0.75–1.25 m), the engineering problems solved, and the alignment of the "
           f"work with the UN Sustainable Development Goals.")
    R.para("**Keywords:** RF fingerprinting, physical-layer authentication, OOK/ASK, EV1527, ESP32, TinyML, "
           "Mahalanobis distance, open-set recognition, replay attack, two-factor smart lock.", space_before=6)

    # -------------------------------------------------------------- contents
    R.h1("Contents", numbered=False, toc=False)
    R.toc_block(prev["toc"], pages)
    R.h1("List of Figures", numbered=False)
    for f in prev["figs"]:
        R.toc_block([(3, f)], pages)
    R.h1("List of Tables", numbered=False, new_page=False)
    for t in prev["tabs"]:
        R.toc_block([(3, t)], pages)

    R.h1("List of Abbreviations", numbered=False)
    R.table(["Abbreviation", "Meaning"], [
        ["ASK / OOK", "Amplitude-shift keying / on–off keying"],
        ["AUC", "Area under the ROC curve"],
        ["CE", "Cross-entropy loss"],
        ["EER", "Equal error rate (FAR = FRR)"],
        ["EV1527", "One-time-programmable 20-bit-address remote-control encoder IC"],
        ["FAR / FRR", "False accept rate / false reject rate"],
        ["FSM", "Finite-state machine"],
        ["ISR", "Interrupt service routine"],
        ["LW", "Ledoit–Wolf covariance shrinkage"],
        ["MAC", "Multiply–accumulate operation"],
        ["MLP", "Multi-layer perceptron"],
        ["NVS", "Non-volatile storage (ESP32 flash key–value store)"],
        ["PLA / RFFI", "Physical-layer authentication / RF fingerprint identification"],
        ["RMT", "ESP32 remote-control (pulse) peripheral"],
        ["ROC", "Receiver operating characteristic"],
        ["SDG", "UN Sustainable Development Goal"],
        ["SDR", "Software-defined radio"],
        ["T", "Elementary symbol period of the encoder (about 335 µs for FOB_A)"],
        ["TinyML", "Machine learning on microcontrollers"],
    ], "Abbreviations and symbols", widths=[4.0, 12.4], align="ll", size=10.5, bold_first=True, numbered=False)

    # =================================================================== CH 1
    R.h1("Introduction & Literature Survey")
    R.h2("Background")
    R.para("Sub-GHz remote controls are among the most widely deployed wireless devices. A typical 433.92 MHz key fob "
           "contains a single encoder chip (EV1527, PT2262 or HT12E), a SAW-stabilised oscillator and a coin cell. "
           "When a button is pressed, the encoder emits a short burst of identical frames by switching the carrier on "
           "and off (OOK). The receiver, often a $1 superheterodyne module, slices the envelope back into a logic "
           "waveform, decodes the 24-bit code and drives a relay if the code is known.")
    R.para("Security therefore rests entirely on the secrecy of a 20-bit address that is broadcast in clear every "
           "time the button is pressed. Capturing it needs nothing more than a receiver of the same kind; replaying it "
           "needs a transmitter module costing less than the remote itself. Even rolling-code systems have been "
           "defeated by jam-and-replay attacks [10], [11]. The root cause is architectural: the receiver checks "
           "*what* was sent, but not *who* sent it.")
    R.para("Every physical transmitter, however, imprints small analogue imperfections on its signal: the exact "
           "frequency of its RC or crystal timebase, its jitter, the rise and fall behaviour of its power amplifier, "
           "and the way its oscillator settles after power-up. These imperfections are not part of the code and are "
           "hard to copy with a different circuit. Measuring them to identify the device is called **RF fingerprinting** "
           "or **physical-layer authentication** [4], [9], [27].")
    R.para("Digital Communication Systems (BECE306P) treats the transmitter as an ideal modulator and the channel "
           "as the only source of impairment [21]. This project turns that view around: the *transmitter's* impairments "
           "become the signal of interest, and the receiver's own distortion becomes the main obstacle to measuring "
           "them. It combines line coding (EV1527 pulse-width coding), synchronisation (symbol-clock recovery), "
           "detection theory (one-class hypothesis testing with FAR/FRR trade-offs) and embedded implementation.")

    R.h2("Literature Survey")
    R.para("Table 1.1 summarises the work that most directly shaped this project. Early research showed that "
           "devices can be identified remotely from clock skew [1], [28] and from the turn-on transient of the "
           "carrier [3], [29]. PARADIS [2] used modulation-domain errors of 802.11 cards measured with a vector "
           "signal analyser. Deep-learning approaches since 2018 learn the fingerprint directly from raw I/Q samples "
           "[5]–[8]; surveys [4], [9], [27] note that nearly all of this work uses laboratory SDRs with high-rate ADCs.")
    R.table(["Ref.", "Signal / platform", "Method", "Relevance and limitation for this project"], [
        ["[1] Kohno 2005", "TCP timestamps over the Internet", "Clock-skew estimation", "Timebase error identifies a device; needs long observation"],
        ["[28] Jana 2010", "802.11 beacons", "Clock skew of access points", "Skew detects rogue APs; needs synchronised timestamps"],
        ["[3] Danev 2009", "IEEE 802.15.4 nodes, high-rate oscilloscope", "Turn-on transient features", "Very low EER, but GS/s capture equipment"],
        ["[29] Ureten 2007", "Transient of RF bursts", "Envelope transient + classifier", "Shows transients are identifying; analogue front end needed"],
        ["[2] Brik 2008 (PARADIS)", "802.11b NICs, vector signal analyser", "Modulation errors + classifier", "High accuracy; lab-grade instrument"],
        ["[5] Merchant 2018", "ZigBee, SDR I/Q", "CNN on error signal", "Deep learning on I/Q; SDR + GPU"],
        ["[6], [7] Riyaz, Sankhe", "Wi-Fi radios, USRP SDR", "CNN (ORACLE)", "Robust to channel; costly SDR testbed"],
        ["[8] Shen 2022", "LoRa, SDR", "Spectrogram CNN, channel-robust", "Addresses channel drift; SDR required"],
        ["[10] Garcia 2016", "Automotive RKE", "Cryptanalysis + replay", "Motivates hardware-level authentication"],
        ["[14] Lee 2018", "Deep features", "Mahalanobis OOD detection", "Basis for our open-set (unknown-device) rejection"],
        ["[17], [18] TinyML", "Microcontrollers", "Quantised small networks", "Feasibility of on-device inference"],
    ], "Summary of related work", widths=[2.8, 3.8, 3.6, 6.2], align="llll", size=9.5)
    R.para("Three observations follow from the survey. First, timing-domain fingerprints (clock skew, jitter, "
           "transient timing) are physically meaningful and survive envelope detection, so they can be measured "
           "*without* I/Q sampling. Second, open-set rejection is essential: a closed-set classifier always assigns an "
           "unknown or cloned transmitter to one of the enrolled classes; distance-based methods [13], [14], [26] with "
           "metric-learning embeddings [15], [16] avoid this. Third, microcontroller-class inference is mature [17], "
           "[18], but we did not find RF fingerprinting demonstrated end-to-end on a $5 MCU with a $1 OOK receiver.")

    R.h2("Research Gap")
    R.bullets([
        "Existing RFFI systems rely on SDRs or oscilloscopes; the receiver itself usually costs more than 100 times the device it protects.",
        "Cheap superheterodyne OOK receivers heavily distort pulse widths (we measured up to 0.33 T); published decoders and fingerprinting methods do not account for this.",
        "Few works close the loop from capture to an *on-device* accept/reject decision that drives an actuator, with reproducible training and deployment.",
        "Physical fingerprints are usually studied as a stand-alone factor; combining them with a user secret on the same unmodified remote has not been explored for low-cost access control.",
    ])

    R.h2("Project Objectives")
    R.bullets([
        "Decode EV1527 frames reliably from a real QC5681 remote through the distortion of an SRX882 receiver, at a usable distance.",
        "Extract code-independent timing features that identify the physical transmitter, with nanosecond-resolution timestamps on an ESP32.",
        "Train an embedding model and an open-set authenticator that accepts the enrolled remote and rejects (a) a transmitter replaying its exact code and (b) unknown transmitters.",
        "Run the authenticator on the ESP32 itself (TinyML, < 8 KB) and keep it bit-exact with the Python implementation.",
        "Build a two-factor smart lock on the receiver: RF fingerprint (have) + secret ON/OFF press pattern (know), with lockout and status indication (LED, buzzer).",
        "Provide a full-stack tool (recorder, trainer, exporter, live dashboard) and evaluate the system honestly on real hardware, including its limits.",
    ], numbered=True)
    R.h2("Scope and Organisation of the Report")
    R.para("The work is limited to the hardware listed in Chapter 2: one genuine remote (two buttons), one receiver "
           "node and one ESP32-S3 + STX882 node used to emulate further transmitters and attacks. Chapter 2 develops "
           "the framework and methodology, Chapter 3 the innovations and feasibility, Chapter 4 the results, Chapter 5 "
           "the SDG alignment and Chapter 6 concludes. Appendices give the pin map, commands and the software layout.")

    # =================================================================== CH 2
    R.h1("Conceptual Framework & Methodology")
    R.h2("Problem Formulation")
    R.para("A press of a transmitter produces a burst of frames x_{1}, x_{2}, …. Each frame carries a code "
           "that *claims* an identity c (the enrolled device that owns that code). The receiver must decide between")
    R.equation(r"H_0:\ \mathbf{x} \sim \mathrm{device}\ c \qquad \mathrm{versus} \qquad H_1:\ \mathbf{x} \sim \mathrm{any\ other\ transmitter}")
    R.para("without ever seeing examples of H_{1} during training (an attacker's hardware is unknown in advance). "
           "This is a one-class, open-set verification problem. The design goals are a false-accept rate of zero "
           "for code clones and unknown devices, a genuine false-reject rate of a few percent per press, a decision "
           "within a fraction of a second, and computation that fits a microcontroller. A second problem, the smart "
           "lock, is a sequential decision over presses: open only if the pattern is correct *and* every press passed H_{0}.")

    R.h2("Signal Model: EV1527 On–Off Keying")
    R.para("The EV1527 encodes 20 address bits and 4 key bits (24 bits). Every bit lasts 4T: a '0' is 1T HIGH + 3T "
           "LOW and a '1' is 3T HIGH + 1T LOW. Each frame ends with a sync symbol of 1T HIGH followed by 31T LOW, "
           "and frames repeat while the button is held [24]. For the QC5681 remote the ON button gives code 0x936768 "
           "(key nibble 8) and OFF gives 0x936764 (key nibble 4); the measured symbol period is T ≈ 335 µs, so")
    R.equation(r"T_{frame} = 24 \cdot 4T + 32T = 128\,T \approx 42.9\ \mathrm{ms}")
    R.para("The SRX882 is a superheterodyne receiver with an envelope detector and a data slicer. Its slicer "
           "threshold lags the envelope, so every HIGH is shortened and every LOW lengthened by an amount δ "
           "that depends on signal strength. With per-pulse noise εᵢ the received widths are")
    R.equation(r"w_i^{H} = u_i\,T - \delta + \varepsilon_i, \qquad w_i^{L} = u_i\,T + \delta + \varepsilon_i, \qquad u_i \in \{1, 3, 31\}")
    R.para(f"{R.next_fig()} shows a real frame of the genuine remote: every HIGH is 10–17 % short and every LOW "
           f"10–17 % long. Close to the antenna δ reaches 0.33T, and then a 3:1 / 1:3 pattern at T fits a 4:2 / 1:5 "
           f"grid at 2T/3 almost equally well, which made the first, conventional decoder lock onto the wrong clock.")
    R.figure(FIG / "fig_real_frame.png", "A real FOB_A frame captured by the receiver at 1 m (top) and the per-pulse "
             "timing residual relative to the recovered grid (bottom). The systematic HIGH/LOW bias is receiver "
             "distortion; the small pulse-to-pulse variation on top of it carries the fingerprint.", 15.0)
    R.figure(FIG / "fig_width_hist.png", "Distribution of all FOB_A pulse widths in units of the recovered clock T "
             "(about 20 000 pulses). HIGH and LOW symbols form separate, shifted clusters around 1T and 3T.", 14.5)

    R.h2("System Architecture")
    R.para(f"{R.next_fig()} shows the complete system. Transmitters (the genuine remote and the ESP32-S3 emulator) "
           "send OOK at 433.92 MHz. The receiver node captures edges, decodes, extracts features and authenticates "
           "on-chip; the same frames are streamed as JSON lines to the laptop, where they are recorded, used for "
           "training and decided a second time by the larger host model. The trained edge model is exported as C "
           "arrays and flashed back to the ESP32, closing the loop.")
    R.figure(FIG / "fig_architecture.png", "System architecture of PhyAuth-RF: transmitters, the receiver/verifier "
             "node (all decisions on-chip) and the laptop tool-chain (record, train, export, visualise).", 16.4)

    R.h2("Hardware Design")
    R.para("Only two ESP32 boards and one module pair were available, so the hardware was split into a "
           "receiver/verifier node and a transmitter/emulator node (Table 2.1). The receiver ESP32 adds a green "
           "status LED with a 470 Ω series resistor and a piezo buzzer for the lock.")
    R.table(["Item", "Part", "Role / connection"], [
        ["Receiver MCU", "ESP32 DevKit (ESP32-D0WD-V3, 240 MHz, CP2102 USB-UART)", "Edge timestamping, DSP, TinyML, lock FSM; USB to laptop (921 600 baud)"],
        ["Receiver RF", "SRX882 superheterodyne ASK receiver, 433.92 MHz", "VCC, CS → 3V3; GND; DATA → GPIO14"],
        ["Antennas", "Helical spring antennas (433 MHz)", "Soldered to SRX882 and STX882 ANT pads"],
        ["Status LED", "Green 5 mm LED + 470 Ω", "GPIO25 → resistor → LED → GND"],
        ["Buzzer", "Piezo buzzer (5 V type)", "GPIO27 → buzzer → GND (quiet at 3.3 V; NPN driver recommended)"],
        ["Transmitter MCU", "ESP32-S3-DevKitC-1 (N16R8)", "RMT pulse generation (100 ns ticks), BOOT button = manual press"],
        ["Transmitter RF", "STX882 ASK transmitter, 433.92 MHz", "DATA ← GPIO4; VCC → 3V3; GND"],
        ["Genuine device", "QC5681 two-button remote, EV1527, 23A 12 V cell", "FOB_A: ON 0x936768, OFF 0x936764"],
    ], "Hardware used", widths=[2.8, 6.2, 7.4], align="lll", size=9.5, bold_first=True)
    R.figure(FIG / "fig_receiver_labelled.png", "Receiver / verifier node (labelled): ESP32 DevKit, SRX882 receiver "
             "with helical antenna, green status LED with 470 Ω resistor, and piezo buzzer.", 15.5)
    R.figure(FIG / "fig_transmitter_labelled.png", "Transmitter / emulator node (labelled): ESP32-S3-DevKitC (N16R8) "
             "driving an STX882 ASK transmitter with a helical antenna.", 14.0)
    R.figure(FIG / "fig_remote_labelled.png", "Genuine device FOB_A: QC5681 two-button EV1527 remote (front view).", 10.0)
    R.figure(FIG / "fig_wiring.png", "Circuit schematic of the two nodes. The SRX882 is powered from 3.3 V so that its DATA "
             "output is a safe logic level for the ESP32.", 16.4)

    R.h2("Receiver Firmware: Nanosecond Edge Capture")
    R.para("A GPIO interrupt on any edge of GPIO14 reads the CPU cycle counter [22] (240 MHz, 4.17 ns per count) and "
           "pushes it into a 1024-entry lock-free ring buffer; nothing else happens in the ISR. The main loop "
           "converts edges into pulse widths, merges glitches shorter than 80 µs and cuts a frame when the "
           "line stays LOW for 4 ms, keeping the preceding gap g (the sync LOW) with the frame. Frames with fewer "
           "than 16 pulses are discarded as noise. Each frame is printed as one JSON line (protocol v2), "
           "together with periodic status lines used by the laptop as a link-health heartbeat:")
    R.code('{"type":"frame","seq":1,"t_us":3921291,"g":10377700,"n":49,"d":[929400,419500,258450,1064275,...]}\n'
           '{"type":"auth","verdict":"AUTHENTIC","dev":"FOB_A","code":"936768/24","score":0.81,"frames":4,...}\n'
           '{"type":"lock","event":"progress","state":"LOCKED","n":3,"len":9,"fails":0,"led":0}')
    R.para("The transmitter node generates frames with the ESP32-S3 RMT peripheral (100 ns ticks) [23]. An important "
           "detail found during testing: sending each frame as a separate RMT call left a 2 ms gap between frames "
           "that corrupted the last bit; the firmware therefore builds the whole press (all frames and the closing "
           "sync HIGH) as one continuous waveform.")

    R.h2("Pulse DSP: Distortion-Invariant Clock Recovery and Decoding")
    R.para("Two quantities are independent of the slicer bias δ because δ cancels inside a HIGH–LOW pair: "
           "the pair period P (which is 4T for EV1527) and the difference between long and short HIGHs. The sync "
           "gap, 31T, then fixes the number of units per pair uniquely:")
    R.equation(r"P_k = w_{2k}^{H} + w_{2k+1}^{L} = nT, \qquad n^{*} = \{\, n \in 3..8 : |g\,n/\bar{P} - 31| \leq 0.08 \cdot 31 \,\}, \qquad \hat{T} = \bar{P}/n^{*}")
    R.para("If the gap is unavailable, n is estimated from the long/short HIGH difference, n = 2P/(P − ΔH). "
           "Bits are then decided per pair (\"the HIGH is the long half\"), which is again immune to δ. A "
           "quality gate rejects frames whose pair-period RMS error exceeds 8 % of T or whose pulses lie more "
           "than 0.45T from the grid. For every accepted frame the residual of each pulse is")
    R.equation(r"r_i = \frac{w_i - u_i \hat{T}}{\hat{T}}, \qquad u_i = \mathrm{round}(w_i/\hat{T})")
    R.figure(FIG / "fig_pipeline.png", "Flowchart of the processing chain from a radio edge to an authentication decision (identical "
             "in the ESP32 C code and in Python).", 15.6)

    R.h2("Physical Fingerprint Features")
    R.para("Eleven features are computed from the residuals and timing of each frame (Table 2.2). None of them "
           "depends on the transmitted bits, so the same model works for every button and code of a device.")
    R.table(["#", "Feature", "Definition", "Physical origin"], [
        ["0", "clock_T_us", "Recovered symbol period T̂", "Encoder RC/crystal timebase frequency"],
        ["1", "pair_jitter_ppm", "Std. of bit-period error (ppm)", "Oscillator cycle-to-cycle jitter"],
        ["2", "duty_short_high", "Mean (H − T)/T, short HIGHs", "PA turn-on/off asymmetry + slicer"],
        ["3", "duty_long_high", "Mean (H − 3T)/T, long HIGHs", "Edge asymmetry at long symbols"],
        ["4", "high_jitter", "Std. of HIGH residuals", "Amplitude/edge noise of the PA"],
        ["5", "first_edge_resid", "Residual of the first HIGH", "Turn-on transient of the transmitter"],
        ["6", "drift_ppm_per_bit", "Slope of bit-period error vs bit index", "Oscillator warm-up chirp"],
        ["7", "sync_ratio", "Preceding sync gap g / T̂", "Encoder sync timing"],
        ["8", "tail_high_resid", "Residual of the trailing sync HIGH", "Turn-off behaviour"],
        ["9", "duty_short_low", "Mean (L − T)/T, short LOWs", "Complementary edge asymmetry"],
        ["10", "duty_long_low", "Mean (L − 3T)/T, long LOWs", "Complementary edge asymmetry"],
    ], "The eleven code-invariant timing features", widths=[0.8, 3.6, 5.6, 6.4], align="clll", size=9.5)

    R.h2("Learning: Embedding Models")
    R.para("Two models are trained on the same frames. The **host model** (FingerprintNet, 21 427 parameters) fuses "
           "a 1D ResNet [19] over the per-bit timing-error sequence (3 × 32) with an MLP over the 11 features. The "
           "**edge model** (EdgeMLP, 963 parameters) uses only the features, so it can run on the ESP32:")
    R.equation(r"z = \frac{x - \mu}{\sigma}, \qquad e = \frac{W_2\,\mathrm{ReLU}(W_1 z + b_1) + b_2}{\| W_2\,\mathrm{ReLU}(W_1 z + b_1) + b_2 \|_2} \in \mathbb{R}^{16}")
    R.para("Both are trained with label-smoothed cross-entropy plus a batch-hard triplet loss [15], [16] that pulls "
           "frames of one device together and pushes other devices at least a margin m = 0.3 away on the unit sphere:")
    R.equation(r"\mathcal{L} = \mathrm{CE}_{0.05}(y, 8\,W_c e) + \frac{1}{B}\sum_a \max\left(0,\ \max_{p} \|e_a - e_p\|^2 - \min_{n} \|e_a - e_n\|^2 + 0.3\right)")
    R.para("Training uses AdamW [20] with cosine learning-rate decay for 60 epochs and a class-balanced sampler. "
           "Augmentation imitates the variation seen between days and distances: T scaled by N(1, 0.008), HIGH "
           "jitter scaled by exp(N(0, 0.35)) and drift perturbed by N(0, 150) ppm/bit. Splits are *temporal* "
           "(later presses are held out), never random, so test frames come from presses the model has not seen. "
           "The OFF button of the remote is recorded as a *variant* (FOB_A_OFF) and merged into FOB_A for training, "
           "so the fingerprint is learned per transmitter rather than per button.")

    R.h2("Open-Set Authenticator and Burst Fusion")
    R.para("For each enrolled device c the centroid and covariance of its validation embeddings are "
           "estimated. Because only tens to hundreds of frames are available in 16 dimensions, the covariance is "
           "shrunk towards a scaled identity (Ledoit–Wolf [12]):")
    R.equation(r"\Sigma_c = (1-\lambda)\,S_c + \lambda\,\frac{\mathrm{tr}(S_c)}{d}\,I, \qquad d_c(e) = \sqrt{(e-\mu_c)^{T}\,\Sigma_c^{-1}\,(e-\mu_c)}")
    R.para("The threshold is set from genuine data only (one-class), as the 97th percentile of genuine validation "
           "distances times a drift margin of 1.3; attacker data are never used for tuning, so reported attack "
           "results are not optimistic. Each frame gets a normalised score and a press is decided on the median "
           "of its first four frames:")
    R.equation(r"\tau_c = 1.3 \cdot Q_{0.97}\left(d_c^{\,genuine}\right), \qquad s_j = \frac{d_c(e_j)}{\tau_c}, \qquad \mathrm{AUTHENTIC} \Leftrightarrow \mathrm{median}(s_1,\ldots,s_4) \leq 1")
    R.para("If the decoded code belongs to an enrolled device and the score fails, the verdict is SPOOF SUSPECTED "
           "(code clone). If the code is unknown, the nearest profile is reported and the verdict is UNKNOWN DEVICE. "
           "A press with fewer than three good frames is not decided.")
    R.para("**Liveness (start-up drift).** A remote is unpowered until pressed, so its oscillator settles during "
           "the first frames; an always-on crystal transmitter shows no such drift. The per-press drift")
    R.equation(r"\rho = \frac{\hat{T}_{first} - \hat{T}_{last}}{\hat{T}_{last}} \cdot 10^{6}\ \mathrm{ppm}")
    R.para("is profiled per device, and the check is switched on only if the accept window (median ± 5 sd) "
           "excludes zero by at least 300 ppm. Otherwise it stays off rather than rejecting genuine presses (see Section 4.7).")

    R.h2("Edge Deployment")
    R.para("The exporter writes the scaler, EdgeMLP weights, per-device centroids, inverse covariances, thresholds, "
           "codes and liveness profiles into `phyauth_model.h` (7 012 bytes of constant arrays). The C implementation "
           "of the DSP and the model (`phyauth_dsp.c`) mirrors the Python code line by line; a parity test compiles "
           "it on the host and checks that both produce the same features, scores and verdicts on recorded frames. "
           "Per frame the on-chip model needs about 1.1 k multiply–accumulates (11×32 + 32×16 + 16×16), "
           "negligible beside the 42.9 ms frame time. The firmware builds for both the ESP32 and the ESP32-S3.")

    R.h2("Two-Factor Smart Lock")
    R.para(f"The lock is a pure C state machine (`phyauth_lock.c`) driven by decided presses ({R.next_fig()}). Factor "
           "1 is the RF fingerprint: a press counts as genuine only if the authenticator accepts it as the lock's "
           "owner. Factor 2 is a secret sequence of ON/OFF presses, 4–12 long (the demonstration used 9). Design rules:")
    R.bullets([
        "Every press advances the attempt; the result is judged only at full length, so an observer cannot learn which press was wrong.",
        "A press rejected by the fingerprint (clone, unknown or unclear) is counted but spoils the attempt: a clone can never unlock even with the right pattern.",
        "Presses from another enrolled device are ignored; 30 s of inactivity abandons an unfinished attempt without penalty.",
        "Three consecutive failures cause a lockout of 30 s that doubles with every further failure (up to 8×).",
        "After unlocking, entering the pattern again locks; auto-relock after 15 s (configurable) or on command.",
        "Learn mode: the owner enters the new pattern twice on the remote; it is stored in NVS flash and survives reboots and model updates.",
        "Outputs: green LED on while unlocked, short flash per accepted press, slow blink in learn mode, fast blink with buzzer on failure or lockout; a relay output (GPIO33) follows the unlocked state.",
    ])
    R.figure(FIG / "fig_lock_fsm.png", "State machine of the two-factor pattern lock running on the receiver ESP32.", 16.4)

    R.h2("Software Stack and Dashboard")
    R.para("The laptop runs a FastAPI back-end (serial link, recorder, trainer, exporter, SQLite event log, REST + "
           "WebSocket) and a Next.js dashboard with two workspaces: **RF Fingerprinting** (Live, Capture, Model lab & "
           "deploy, Attack lab, Devices, Hardware) and **Smart Lock** (padlock view, two factor cards, timeline, settings). "
           "The serial reader has a watchdog: the receiver sends a status line every 2 s, so 8 s of silence reopens "
           "the port automatically (this fixed a stall after re-flashing). In total about 10 600 lines of Python, "
           "C/C++ and TypeScript were written, with 35 automated tests.")

    R.h2("Experimental Methodology")
    R.para(f"{R.next_tab()} lists the transmitters. Recordings were made at 0.75 m, 1.0 m and 1.25 m on three sessions "
           "(s01 on 17 Oct, s02/s03 on 21 Oct 2026), holding each button for a normal press.")
    R.table(["ID", "Hardware", "Role", "Code", "Notes"], [
        ["FOB_A", "QC5681 remote, ON button", "Genuine, enrolled", "936768/24", "T ≈ 335 µs"],
        ["FOB_A_OFF", "QC5681 remote, OFF button", "Variant of FOB_A", "936764/24", "Merged into FOB_A"],
        ["EMU_1", "ESP32-S3 + STX882, profile 1", "Enrolled (emulated)", "5a3c91/24", "T 318 µs, 2500 ppm jitter"],
        ["EMU_2", "ESP32-S3 + STX882, profile 2", "Enrolled (emulated)", "1e77c4/24", "T 331 µs, 3500 ppm jitter"],
        ["EMU_3", "ESP32-S3 + STX882, profile 3", "Unknown transmitter", "9b0d2a/24", "Never enrolled"],
        ["CLONE_A", "ESP32-S3 + STX882, clean", "Attacker: replays FOB_A", "936768/24", "Same code and T as FOB_A"],
    ], "Transmitters used in the experiments", widths=[2.2, 4.4, 3.6, 2.4, 3.8], align="lllll", size=9.5, bold_first=True)
    R.para("Metrics: FAR = accepted impostor decisions / impostor decisions; FRR = rejected genuine decisions / "
           "genuine decisions; EER is the error rate where FAR = FRR when the threshold is swept; AUC is the area "
           "under the ROC curve. They are reported per frame and per burst (press). Attack rejection is the fraction "
           "of clone or unknown bursts that are not accepted.")
    R.para("**Honest labelling of emulated devices.** EMU_1 and EMU_2 differ from each other by *programmed* "
           "timing (period, jitter, skew), which is a weaker difference than between two real remotes. Results that "
           "separate EMU_1 from EMU_2 are therefore a functional test of the pipeline; the security-relevant results "
           "are those for the genuine remote against CLONE_A and EMU_3, which are physically different transmitters.")

    R.h2("Implementation Plan")
    R.table(["Phase", "Work", "Outcome"], [
        ["1", "Firmware for edge capture and RMT transmission; serial protocol", "Raw frames from both boards"],
        ["2", "Bench tests of distance, antenna, receiver distortion", "0.75–1.25 m window; distortion model"],
        ["3", "Distortion-invariant clock and pair decoder; regression tests", "Remote decoded at all distances"],
        ["4", "Features, models, open-set authenticator, evaluation", "Model + metrics report"],
        ["5", "C port, exporter, parity test, on-chip decisions", "Authentication on the ESP32"],
        ["6", "Two-factor lock, LED/buzzer, learn mode, flash storage", "Working lock prototype"],
        ["7", "Dashboard (two workspaces), watchdog, data collection, retraining", "Final system and this report"],
    ], "Implementation plan", widths=[1.4, 9.0, 6.0], align="cll", size=10)

    # =================================================================== CH 3
    R.h1("Innovation & Feasible Solution")
    R.h2("Novel Aspects of the Solution")
    R.bullets([
        "**RF fingerprinting without an ADC or SDR.** All features come from edge times measured with the CPU cycle counter (4.17 ns) behind a $1 envelope receiver.",
        "**Distortion-invariant clock recovery.** Using the pair period and the 31T sync gap removes the receiver's slicer bias; the conventional histogram method locked onto 2T/3 and decoded nothing from the real remote.",
        "**Code-invariant fingerprint.** Features never use the bit values, so one model covers both buttons; button variants are merged per transmitter.",
        "**On-device open-set authentication.** A 963-parameter network plus Mahalanobis profiles (7 KB) decides each press on the ESP32 and rejects transmitters it has never seen; Python and C are kept identical by a parity test.",
        "**Two factors on one unmodified remote.** The smart lock combines the physical fingerprint (have) with a press pattern (know); a cloned code with the correct pattern still fails, and the pattern is never revealed press by press.",
        "**Self-validating liveness.** The start-up drift check enables itself only when it is statistically safe, avoiding false rejections.",
        "**Complete reproducible tool-chain.** Record → train → evaluate → export → flash → live dashboard, with 35 automated tests and an emulator for controlled attacks.",
    ])
    R.h2("Comparison with Conventional Approaches")
    R.table(["Aspect", "Typical SDR-based RFFI", "Fixed-code receiver", "PhyAuth-RF"], [
        ["Receiver front end", "SDR with 10–100 MS/s I/Q ADC", "OOK module", "OOK module (SRX882)"],
        ["Measured quantity", "Raw I/Q waveform", "Bits only", "Edge times at 4.17 ns"],
        ["Processing", "PC/GPU", "MCU decoder", "ESP32 (on-chip) + optional PC"],
        ["Detects code clone", "Yes", "No", "Yes (30/30 bursts)"],
        ["Rejects unknown device", "Depends on method", "No", "Yes, open-set (30/30 bursts)"],
        ["User secret factor", "No", "No", "Yes, ON/OFF pattern"],
        ["Approx. receiver cost", "Hundreds to thousands of USD", "< Rs. 200", "≈ Rs. 650"],
    ], "Qualitative comparison", widths=[3.6, 4.4, 3.4, 5.0], align="llll", size=9.5, bold_first=True)

    R.h2("Technical Feasibility")
    R.bullets([
        "**Timing resolution:** 4.17 ns per count vs. a symbol of 335 µs, i.e. 1 part in 80 000; the fingerprint features vary at the 10^{4} ppm level and are easily resolved.",
        "**Computation:** about 1.1 k MACs and an O(pulses) DSP pass per frame on a 240 MHz core; the decision latency (≈ 172 ms for four frames) is set by the air time of the frames, not by computation.",
        "**Memory:** model header 7 012 bytes; ring buffer 4 KB; lock state < 100 bytes; pattern in NVS flash.",
        "**Interfaces:** a single GPIO input for the receiver; LED, buzzer and relay outputs; USB serial only for monitoring and training. The node works stand-alone once flashed.",
        "**Range:** a reliable decoding window of 0.75–1.25 m with the spring antennas (Section 4.5); enough for a door-side reader, and extendable with a better antenna or a receiver with AGC.",
    ])
    R.h2("Economic Feasibility")
    R.table(["Component", "Qty", "Approx. cost (Rs.)"], [
        ["ESP32 DevKit (receiver)", "1", "400"],
        ["ESP32-S3-DevKitC N16R8 (transmitter / emulator)", "1", "800"],
        ["SRX882 + STX882 433 MHz module pair with spring antennas", "1 set", "300"],
        ["QC5681 two-button EV1527 remote", "1", "250"],
        ["Breadboards, jumper wires", "2 + set", "240"],
        ["Green LED, 470 Ω resistor, piezo buzzer", "1 each", "30"],
        ["**Total development kit**", "", "**≈ 2 020**"],
        ["**Deployable verifier only** (ESP32 + SRX882 + LED/buzzer)", "", "**≈ 600–650**"],
    ], "Bill of materials (approximate Indian retail prices, October 2026)", widths=[10.0, 2.4, 4.0], align="lcr", size=10)
    R.para("A deployable verifier costs about as much as the receiver it would replace plus a microcontroller, and "
           "protects existing remotes without changing them, so no new key fobs need to be bought or discarded.")
    R.h2("Safety Considerations")
    R.bullets([
        "**Radio:** 433.92 MHz lies in the license-exempt short-range-device band (433–434.79 MHz in India); the STX882 is a low-power module and transmissions are short bursts.",
        "**Electrical:** all logic is 3.3 V from USB; the SRX882 is powered from 3.3 V so its DATA pin cannot exceed the ESP32's input limit; the LED current is limited by 470 Ω (≈ 3 mA).",
        "**Actuators:** a relay or lock driver must be electrically isolated and should fail secure; mains loads must not be wired on a breadboard.",
        "**Buzzer:** the 5 V buzzer draws more than a GPIO should source at full volume; an NPN transistor driver is required for production.",
        "**Responsible testing:** the clone emulator replays only our own remote's code in the lab; no third-party systems were targeted.",
    ])
    R.h2("Operational Constraints")
    R.bullets([
        "Each remote must be enrolled (about 30–40 presses per button and distance) before use, and a new receiver needs new enrolment, because the receiver's distortion is part of every measurement.",
        "Fingerprints drift between days, temperatures and battery levels; recording sessions on several days and augmentation mitigate this, and re-enrolment may be needed after a battery change.",
        "Decoding requires 0.75–1.25 m with the current antennas; strong signals saturate the slicer and weak ones lose short pulses.",
    ])

    # =================================================================== CH 4
    R.h1("Results & Discussion")
    R.h2("Physical Prototype")
    R.para("Both nodes were built on breadboards (Figures 2.4–2.6) and run the final firmware v2.1. The final "
           f"deployed model is version 20261021-224341 (trained in {REP['meta']['train_seconds']} s on the laptop), "
           "flashed to the receiver ESP32; the laptop runs the same version, so both deciders agree.")
    R.h2("Dataset")
    R.para("Of 1 602 recorded frames, 1 442 (90.0 %) passed the quality gate. Most rejections were frames with too "
           "few pulses (typically fragments at the start or end of a press). After removing frames with "
           "a foreign code, training used 712 frames, validation 137 and testing 152 genuine frames, plus 420 "
           "impostor frames (CLONE_A and EMU_3) that were never used for training or threshold selection.")
    R.figure(FIG / "fig_dataset.png", "Frames that passed the quality gate, by transmitter and distance (FOB_A_OFF "
             "includes 131 frames recorded on 21 Oct to strengthen the OFF button).", 13.5)
    R.h2("Feature Space")
    R.para(f"{R.next_fig()} plots two pairs of raw features. Clock period alone separates the emulated profiles and "
           "the unknown transmitter, but **CLONE_A deliberately uses the same code and the same T as FOB_A** and "
           "overlaps it in clock period; it is separated by edge-shape features (duty error, first-edge residual) "
           "that come from its different power amplifier and timebase. The FOB_A clusters spread with distance and "
           "session, which is why multi-distance enrolment matters.")
    R.figure(FIG / "fig_feature_scatter.png", "Raw physical features of every accepted frame. Left: clock period "
             "vs. bit-period jitter (13 of 1 442 frames that locked to 3T lie off-scale). Right: short-HIGH duty error "
             "vs. first-edge residual; the clone (red) is clearly apart from the genuine remote (purple).", 15.3)
    fi = REP["feature_importance"]
    R.table(["Feature shuffled", "Frame EER after shuffling", "Increase in EER"],
            [[f["feature"], pct(f["eer"], 2), f"+{f['delta_eer'] * 100:.2f} %"] for f in fi],
            "Permutation importance of the features for the edge model (test set)", widths=[5.4, 5.0, 4.0],
            align="lcc", size=9.5)
    R.para("Clock period and pair jitter carry most identity, followed by edge-asymmetry features. The sync ratio "
           "contributes nothing for these devices and could be removed to save computation.")
    R.figure(AS / "shot6.png", "Dashboard Model lab: feature importance and PCA of the embedding space (previous "
             "model version 20261021-213950). Enrolled devices form tight clusters; the clone (×, not enrolled) "
             "lands between FOB_A and the emulators, outside every acceptance region.", 14.8)

    R.h2("Authentication Performance")
    R.table(["Metric", "Edge model (on ESP32)", "Host fusion model"], [
        ["Parameters / size", "963 / 7 012 B header", "21 427"],
        ["Closed-set accuracy (152 test frames)", "100 %", "100 %"],
        ["Frame EER / AUC", f"{pct(E['frame']['eer'], 2)} / {E['frame']['auc']:.5f}", f"{pct(F['frame']['eer'], 2)} / {F['frame']['auc']:.4f}"],
        ["Frame FAR / FRR at operating threshold", f"0 % / {pct(E['frame']['frr'])}", f"0 % / {pct(F['frame']['frr'])}"],
        ["Burst EER / AUC", "0 % / 1.0000", "0 % / 1.0000"],
        ["Burst FAR / FRR (27 genuine, 234 impostor bursts)", "0 % / 0 %", "0 % / 0 %"],
        ["Code clone CLONE_A rejected", "30 / 30 bursts (100 %)", "30 / 30 bursts (100 %)"],
        ["Unknown EMU_3 rejected", "30 / 30 bursts (100 %)", "30 / 30 bursts (100 %)"],
        ["Validation EER", pct(REP["meta"]["validation_eer"]["edge"], 2), pct(REP["meta"]["validation_eer"]["fusion"], 2)],
    ], "Results of the final model (version 20261021-224341) on held-out data", widths=[6.6, 5.0, 4.8], align="lcc", size=10, bold_first=True)
    R.para("The edge model is the better *operating* model: the host model separates classes slightly better "
           "(lower frame EER) but its one-class threshold is tighter, so it rejects more genuine single frames "
           f"({pct(F['frame']['frr'])} vs {pct(E['frame']['frr'])}). Fusing frames removes the difference: at the "
           "burst level both make no error on this test set. With only 27 genuine test bursts the true burst FRR "
           "can still be a few percent (a 95 % upper bound of about 11 % for 0 errors in 27), so the live results "
           "below are an important complement.")
    R.figure(AS / "shot5.png", "Dashboard Model lab: ROC on a logarithmic false-accept axis and the closed-set "
             "confusion matrix with per-distance frame FRR (previous model version, 213950).", 14.8)

    R.h2("Distance Robustness")
    lodo = ANA["leave_one_distance_out"]
    rows = []
    for L in lodo:
        g = L["genuine"]
        rows.append([f"{L['held_out_m']} m", *(f"{int(round(g[k]['frr'] * g[k]['presses']))}/{g[k]['presses']}" for k in ("FOB_A", "EMU_1", "EMU_2")),
                     f"{int(L['clone']['rejection'] * L['clone']['presses'])}/{L['clone']['presses']}",
                     f"{int(L['unknown']['rejection'] * L['unknown']['presses'])}/{L['unknown']['presses']}"])
    R.para(f"{R.next_fig()}(a) gives the single-frame FRR of the final models per distance; accuracy was 100 % "
           "everywhere. The bench window in (b) was measured by moving the remote: closer than 0.5 m the receiver "
           "saturates and the slicer bias grows, and from 1.5 m the short 1T pulses are lost. In a leave-one-"
           f"distance-out test ({R.next_tab()}), the model was trained without one distance and tested on it.")
    R.figure(FIG / "fig_distance.png", "(a) Single-frame genuine rejection by distance for both models. "
             "(b) Decoding window of the SRX882 rig observed on the bench.", 15.3)
    R.table(["Held-out distance", "FOB_A rejected", "EMU_1 rejected", "EMU_2 rejected", "Clones rejected", "Unknown rejected"],
            rows, "Leave-one-distance-out test (presses; development model trained on 17 Oct data)", widths=[2.9, 2.6, 2.6, 2.6, 2.8, 2.9], align="cccccc", size=9.5)
    R.para("Only one genuine press out of 90 was rejected when its distance had never been seen in training, "
           "while every clone and unknown press was rejected at every distance.")

    R.h2("Burst Fusion and Latency")
    R.figure(FIG / "fig_frames_decision.png", "Genuine rejection rate and decision latency against the number of "
             "frames fused per press (90 genuine presses).", 12.0)
    R.para("A single frame is a noisy measurement (45.6 % of genuine presses would be rejected), but the median of "
           "four frames brings the FRR to 1.1 % at a latency of 171.6 ms after the first frame, still well within a "
           "normal button press. Clone and unknown rejection stayed at 100 % for every setting, so four frames were "
           "chosen; the firmware decides on three frames if a press is shorter.")

    R.h2("Cross-Session Check and Liveness")
    R.table(["Device", "Presses", "Accepted", "Matched profile", "Median score", "By distance"], [
        ["FOB_A (ON)", str(XDAY["FOB_A"]["presses"]), str(XDAY["FOB_A"]["accepted"]), "FOB_A", f"{XDAY['FOB_A']['median_score']:.2f}",
         ", ".join(f"{k} m: {v}" for k, v in XDAY["FOB_A"]["accepted_by_distance"].items())],
        ["FOB_A_OFF", str(XDAY["FOB_A_OFF"]["presses"]), str(XDAY["FOB_A_OFF"]["accepted"]), "FOB_A", f"{XDAY['FOB_A_OFF']['median_score']:.2f}",
         ", ".join(f"{k} m: {v}" for k, v in XDAY["FOB_A_OFF"]["accepted_by_distance"].items())],
    ], "Presses from a later session (21 Oct) against a model trained on 17 Oct data", widths=[2.6, 1.8, 2.0, 2.6, 2.4, 5.0], align="lccccl", size=9.5)
    R.para("All presses of the genuine remote were accepted on a different day, and OFF presses (a different code) "
           "matched the FOB_A profile without being enrolled separately, confirming that the fingerprint belongs to "
           "the transmitter, not the code. During live lock use, however, OFF presses held differently in the hand "
           "scored close to the threshold (1.0–1.6); 131 further OFF frames were therefore recorded and the model "
           "retrained (final model). Its live re-test is part of the next demonstration.")
    lv = REP["liveness"]
    R.table(["Device", "Presses", "Median drift (ppm)", "Robust sd (ppm)", "Check active?"], [
        [k, str(v["n"]), f"{v['median_ppm']:.0f}", f"{v['sd_ppm']:.0f}", "no – " + v["reason"]] for k, v in lv["profiles"].items()
    ] + [["CLONE_A (impostor)", str(lv["impostor_test"]["CLONE_A"]["presses"]), f"{lv['impostor_test']['CLONE_A']['median_ppm']:.1f}", "–", "–"]],
        "Start-up drift per press (liveness profiles)", widths=[3.4, 1.8, 3.2, 3.0, 5.0], align="lcccl", size=9.5)
    R.para("The genuine remote does show start-up drift (median −1031 ppm) while the always-on clone shows almost "
           "none (5 ppm), which confirms the physical idea. Its spread between presses (sd 526 ppm) is too large for "
           "a safe five-sigma window, so the check correctly stayed off: enabling it would have rejected genuine "
           "presses. A remote with a steadier oscillator, or more frames per press, could enable it.")

    R.h2("Live System and Dashboard")
    R.para(f"{R.next_fig()} and {R.next_fig(2)} show live decisions. A genuine FOB_A press is accepted on the ESP32 "
           "with score 0.91 and on the laptop with 0.82. An EMU_1 press scored 1.08 / 1.06 and was flagged as a "
           "spoof: this is a **false rejection** of a borderline emulated press, consistent with the few-percent "
           f"burst FRR discussed above; the next press was accepted. {R.next_fig(3)} shows the decision log, where "
           "both deciders agree on every press.")
    R.figure(AS / "shot1.png", "Live authentication: genuine remote accepted on the ESP32 (score 0.91, 4 frames, "
             "T = 334.64 µs), with the last frame's pulse train and timing residuals.", 14.8)
    R.figure(AS / "shot2.png", "Live authentication: a borderline EMU_1 press (score 1.08 > 1) flagged as spoof "
             "suspected – an example of a false rejection.", 14.8)
    R.figure(AS / "shot3.png", "Decision log: verdict, device, code, score, frames and which decider (ESP32 or laptop) "
             "made the decision.", 14.8)

    R.h2("Smart Lock Demonstration")
    R.para("The lock was set to a 9-press ON/OFF pattern owned by FOB_A. Entering the pattern on the genuine remote "
           f"unlocked the ESP32 and the dashboard padlock opened ({R.next_fig(2)}); the LED followed the state and the "
           "lock re-armed automatically after 15 s. The lock logic is verified by 10 host unit tests on the same C "
           f"source (Table {R.chapter}.{R.tab_n + 1}).")
    R.figure(AS / "shot7.png", "Smart Lock workspace, locked: padlock, progress dots for the 9-press pattern, the two "
             "factor cards and the settings stored in the ESP32's flash.", 14.8)
    R.figure(AS / "shot8.png", "Smart Lock workspace after the correct pattern on the genuine remote: UNLOCKED, "
             "auto-relock countdown, factor 1 'genuine' (score 0.32).", 14.8)
    R.table(["Test (tests/test_lock.py)", "Checks", "Result"], [
        ["correct_pattern_unlocks_and_relocks", "Right pattern unlocks; same pattern locks again", "pass"],
        ["wrong_pattern_fails_without_revealing_position", "No early failure signal", "pass"],
        ["rejected_press_is_counted_and_fails_once_at_the_end", "Foreign press spoils attempt; one failure", "pass"],
        ["other_enrolled_device_is_ignored", "Other enrolled remotes cannot interfere", "pass"],
        ["three_failures_cause_lockout_that_expires_and_grows", "Lockout 30 s, doubling", "pass"],
        ["abandoned_attempt_is_cleared_without_failure", "Idle timeout without penalty", "pass"],
        ["auto_relock", "Relocks after the configured time", "pass"],
        ["learn_mode_needs_matching_confirmation", "Pattern must be entered twice", "pass"],
        ["clone_presses_can_never_unlock", "Correct pattern from a clone fails", "pass"],
        ["learn_rejects_too_short_and_foreign", "Minimum length; foreign presses ignored", "pass"],
    ], "Unit tests of the lock state machine (all 35 project tests pass)", widths=[7.4, 7.0, 2.0], align="llc", size=9.5)

    R.h2("Comparative Benchmarking")
    R.para("Direct numerical comparison with the literature is limited because the signals, devices and receivers "
           "differ. Qualitatively, laboratory systems [2], [3], [7] report very high accuracy among many devices of "
           "the same model using I/Q or oscilloscope captures; PhyAuth-RF reaches zero burst errors on its test set "
           "with a receiver costing a few hundred rupees and with the full decision on a microcontroller, but for "
           "far fewer devices. Compared with a standard fixed-code receiver, which accepts 100 % of replayed codes, "
           "PhyAuth-RF accepted none of 30 replay bursts.")
    R.h2("Engineering Problems Solved")
    R.table(["Problem observed", "Root cause", "Fix"], [
        ["Real remote decoded 0 frames", "SRX882 slicer bias up to 0.33T; clock locked to 2T/3", "Pair-period + sync-gap clock, pair decoding, regression tests"],
        ["Emulated frames had a wrong last bit", "2 ms RMT latency between per-frame calls", "Whole press as one RMT waveform with closing sync"],
        ["Exported header did not compile", "Float literal written as '1f'", "Exporter writes '1.0f'; parity test"],
        ["Same press both AUTHENTIC and SPOOF", "ESP32 and laptop ran different model versions", "Flash the latest export; version shown in top bar"],
        ["Dashboard froze after re-flash", "Serial reader stuck on a stale port handle", "8 s silence watchdog reopens the port"],
        ["OFF presses near threshold", "Fewer OFF recordings, different grip", "131 more OFF frames, retrained"],
        ["Green LED only blinked", "LED part is a self-flashing type", "Use a plain LED (GPIO verified HIGH)"],
        ["Buzzer silent", "5 V buzzer on a 3.3 V GPIO", "NPN transistor driver (planned)"],
    ], "Problems found during development and their fixes", widths=[4.6, 5.6, 6.2], align="lll", size=9.5)
    R.h2("Limitations and Threats to Validity")
    R.bullets([
        "One genuine remote; the emulated devices differ by programmed timing. A study with several remotes of the same model is needed to quantify inter-device separability.",
        "Small test sets (27 genuine and 60 attack bursts); zero errors does not mean a zero error rate.",
        "The fingerprint is measured through one specific receiver; moving the model to another receiver requires re-enrolment.",
        "Data were collected over two days indoors; temperature and battery ageing were not studied.",
        "A sophisticated attacker with a programmable waveform source could try to imitate the measured timing.",
    ])

    # =================================================================== CH 5
    R.h1("Sustainable Development Goals (SDGs) & Societal Impact")
    R.h2("Alignment with the UN Sustainable Development Goals")
    R.para("The 2030 Agenda [25] defines 17 goals. Table 5.1 maps the project to the most relevant goals and targets.")
    R.table(["SDG", "Relevant target", "Contribution of PhyAuth-RF"], [
        ["SDG 9: Industry, Innovation and Infrastructure", "9.1 resilient infrastructure; 9.5 research and innovation; 9.c ICT access",
         "Low-cost security layer for existing wireless control infrastructure; open, reproducible research tool-chain built on commodity parts"],
        ["SDG 11: Sustainable Cities and Communities", "11.1 safe, adequate housing; 11.7 safe public spaces",
         "Hardens gates, garages and alarm panels of homes and community buildings against replay and cloning"],
        ["SDG 12: Responsible Consumption and Production", "12.5 reduce waste generation",
         "Upgrades security of remotes already in use instead of replacing them with new rolling-code devices, avoiding e-waste"],
        ["SDG 7: Affordable and Clean Energy", "7.3 energy efficiency",
         "Verifier runs on a < 1 W microcontroller instead of a PC + SDR drawing tens of watts; suitable for battery or solar power"],
        ["SDG 4: Quality Education", "4.4 relevant technical skills",
         "A complete, low-cost teaching platform for digital communication, DSP, ML and embedded security"],
    ], "Mapping to SDGs", widths=[4.2, 4.6, 7.6], align="lll", size=9.5, bold_first=True)
    R.h2("Benefits to Society")
    R.bullets([
        "**Safety and security of homes and small businesses:** cheap remotes are used where commercial access-control systems are unaffordable; adding a hardware-identity check closes the simplest attack, code replay.",
        "**Affordability:** about Rs. 650 per protected receiver makes the technique viable in developing regions and rural installations.",
        "**Retrofit, not replace:** users keep their remotes; only the receiver side changes, which also reduces cost and waste.",
        "**Two-factor convenience:** the pattern lock adds a knowledge factor without a keypad, phone or network connection, so it works offline and during network outages.",
        "**Education and awareness:** the Attack lab demonstrates, safely and on one's own devices, why fixed codes are insecure.",
    ])
    R.h2("Environmental and Ethical Considerations")
    R.para("The system uses existing hardware, consumes little power and adds no batteries. Only timing data of "
           "the user's own devices are stored; no personal or location data are collected. The attack emulator is "
           "dual-use and is restricted to the user's own remote in a laboratory; the report deliberately describes "
           "it only at the level needed to evaluate the defence.")

    # =================================================================== CH 6
    R.h1("Conclusion & References")
    R.h2("Conclusion")
    R.para("PhyAuth-RF shows that physical-layer authentication of 433 MHz remotes is possible with a $5 "
           "microcontroller and a $1 receiver. The main outcomes are:")
    R.bullets([
        "Reliable decoding of a real EV1527 remote through severe receiver distortion, using a distortion-invariant clock recovery derived from the line code itself.",
        "Eleven code-invariant timing features measured at 4.17 ns that separate the genuine remote from a transmitter replaying its exact code.",
        f"An on-chip open-set authenticator (963 parameters, 7 KB) with 100 % accuracy, frame EER {pct(E['frame']['eer'], 2)}, burst EER 0 % and 30/30 clone and 30/30 unknown bursts rejected on held-out data.",
        "A two-factor smart lock on the same ESP32 that requires both the genuine remote and a secret ON/OFF pattern, with lockout, learn mode, flash storage and LED/buzzer indication, demonstrated on hardware and covered by unit tests.",
        "A complete open tool-chain (firmware, DSP, training, export, dashboard) with 35 automated tests, and a documented account of the hardware's limits.",
    ])
    R.h2("Future Work")
    R.bullets([
        "Enrol several remotes of the same model and collect data over weeks, temperatures and battery levels.",
        "Add an NPN buzzer driver, a relay-driven lock actuator and an enclosure; power the verifier from a battery.",
        "Better antennas and a receiver with AGC to widen the working range; a second receiver for diversity.",
        "Online adaptation of the profiles from accepted presses to follow slow drift.",
        "Extend the method to other OOK encoders (PT2262, HT12E) and to LoRa/BLE timing.",
    ])
    R.h2("References")
    refs = [
        "T. Kohno, A. Broido, and K. C. Claffy, “Remote physical device fingerprinting,” *IEEE Trans. Dependable Secure Comput.*, vol. 2, no. 2, pp. 93–108, 2005.",
        "V. Brik, S. Banerjee, M. Gruteser, and S. Oh, “Wireless device identification with radiometric signatures,” in *Proc. ACM MobiCom*, 2008, pp. 116–127.",
        "B. Danev and S. Capkun, “Transient-based identification of wireless sensor nodes,” in *Proc. ACM/IEEE IPSN*, 2009, pp. 25–36.",
        "B. Danev, D. Zanetti, and S. Capkun, “On physical-layer identification of wireless devices,” *ACM Comput. Surv.*, vol. 45, no. 1, Art. 6, 2012.",
        "K. Merchant, S. Revay, G. Stantchev, and B. Nousain, “Deep learning for RF device fingerprinting in cognitive communication networks,” *IEEE J. Sel. Topics Signal Process.*, vol. 12, no. 1, pp. 160–167, 2018.",
        "S. Riyaz, K. Sankhe, S. Ioannidis, and K. Chowdhury, “Deep learning convolutional neural networks for radio identification,” *IEEE Commun. Mag.*, vol. 56, no. 9, pp. 146–152, 2018.",
        "K. Sankhe, M. Belgiovine, F. Zhou, S. Riyaz, S. Ioannidis, and K. Chowdhury, “ORACLE: Optimized radio classification through convolutional neural networks,” in *Proc. IEEE INFOCOM*, 2019, pp. 370–378.",
        "G. Shen, J. Zhang, A. Marshall, and J. R. Cavallaro, “Towards scalable and channel-robust radio frequency fingerprint identification for LoRa,” *IEEE Trans. Inf. Forensics Security*, vol. 17, pp. 774–787, 2022.",
        "A. Jagannath, J. Jagannath, and P. S. P. V. Kumar, “A comprehensive survey on radio frequency (RF) fingerprinting: Traditional approaches, deep learning, and open challenges,” *Computer Networks*, vol. 219, Art. 109455, 2022.",
        "F. D. Garcia, D. Oswald, T. Kasper, and P. Pavlidès, “Lock it and still lose it – On the (in)security of automotive remote keyless entry systems,” in *Proc. USENIX Security Symp.*, 2016, pp. 929–944.",
        "S. Kamkar, “Drive it like you hacked it: New attacks and tools to wirelessly steal cars,” presented at DEF CON 23, Las Vegas, NV, USA, 2015.",
        "O. Ledoit and M. Wolf, “A well-conditioned estimator for large-dimensional covariance matrices,” *J. Multivariate Anal.*, vol. 88, no. 2, pp. 365–411, 2004.",
        "P. C. Mahalanobis, “On the generalised distance in statistics,” *Proc. Nat. Inst. Sci. India*, vol. 2, no. 1, pp. 49–55, 1936.",
        "K. Lee, K. Lee, H. Lee, and J. Shin, “A simple unified framework for detecting out-of-distribution samples and adversarial attacks,” in *Proc. NeurIPS*, 2018.",
        "F. Schroff, D. Kalenichenko, and J. Philbin, “FaceNet: A unified embedding for face recognition and clustering,” in *Proc. IEEE CVPR*, 2015, pp. 815–823.",
        "A. Hermans, L. Beyer, and B. Leibe, “In defense of the triplet loss for person re-identification,” arXiv:1703.07737, 2017.",
        "R. David *et al.*, “TensorFlow Lite Micro: Embedded machine learning for TinyML systems,” in *Proc. MLSys*, 2021.",
        "P. Warden and D. Situnayake, *TinyML: Machine Learning with TensorFlow Lite on Arduino and Ultra-Low-Power Microcontrollers*. Sebastopol, CA, USA: O’Reilly, 2019.",
        "K. He, X. Zhang, S. Ren, and J. Sun, “Deep residual learning for image recognition,” in *Proc. IEEE CVPR*, 2016, pp. 770–778.",
        "I. Loshchilov and F. Hutter, “Decoupled weight decay regularization,” in *Proc. ICLR*, 2019.",
        "J. G. Proakis and M. Salehi, *Digital Communications*, 5th ed. New York, NY, USA: McGraw-Hill, 2008.",
        "Espressif Systems, *ESP32 Technical Reference Manual*. Shanghai, China: Espressif Systems.",
        "Espressif Systems, *ESP32-S3 Technical Reference Manual*. Shanghai, China: Espressif Systems.",
        "Silvan Chip Electronics, *EV1527 OTP Encoder Datasheet*.",
        "United Nations General Assembly, “Transforming our world: The 2030 Agenda for Sustainable Development,” A/RES/70/1, 2015.",
        "D. Hendrycks and K. Gimpel, “A baseline for detecting misclassified and out-of-distribution examples in neural networks,” in *Proc. ICLR*, 2017.",
        "Q. Xu, R. Zheng, W. Saad, and Z. Han, “Device fingerprinting in wireless networks: Challenges and opportunities,” *IEEE Commun. Surveys Tuts.*, vol. 18, no. 1, pp. 94–104, 2016.",
        "S. Jana and S. K. Kasera, “On fast and accurate detection of unauthorized wireless access points using clock skews,” *IEEE Trans. Mobile Comput.*, vol. 9, no. 3, pp. 449–462, 2010.",
        "O. Ureten and N. Serinken, “Wireless security through RF fingerprinting,” *Canadian J. Elect. Comput. Eng.*, vol. 32, no. 1, pp. 27–33, 2007.",
    ]
    for i, r in enumerate(refs, 1):
        p = R.para(f"[{i}]\t{r}", align="left", size=10.5, space_after=3)
        p.paragraph_format.left_indent = Cm(1.0)
        p.paragraph_format.first_line_indent = Cm(-1.0)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(1.0))
        p.paragraph_format.line_spacing = 1.05

    # ============================================================== APPENDIX
    R.h1("Appendix A: Pin Map, Commands and Software Organisation", numbered=False)
    R.chapter = "A"
    R.sec = R.tab_n = 0
    R.table(["Node", "Pin", "Function"], [
        ["Receiver ESP32", "GPIO14", "SRX882 DATA input (edge interrupt)"],
        ["Receiver ESP32", "GPIO25", "Green LED (unlocked / accepted press / alarm blink)"],
        ["Receiver ESP32", "GPIO27", "Buzzer (wrong pattern, foreign remote, lockout)"],
        ["Receiver ESP32", "GPIO26 / GPIO33", "Red LED (locked) / relay input (unlocked), optional"],
        ["Receiver ESP32", "GPIO2", "On-board blue LED (verdict indicator)"],
        ["Transmitter ESP32-S3", "GPIO4", "STX882 DATA output (RMT)"],
        ["Transmitter ESP32-S3", "GPIO0 (BOOT)", "Manual press button"],
    ], "Pin map", widths=[4.4, 3.4, 8.6], align="lll", size=10)
    R.table(["Command (receiver)", "Effect"], [
        ["PING / INFO", "Link check / firmware, model version and settings"],
        ["MODE STREAM | EDGE | BOTH", "Stream frames, decide on-chip, or both"],
        ["SET GAP <us> / SET GLITCH <us> / SET MINP <n>", "Framing gap, glitch merge width, minimum pulses"],
        ["LIVE ON | OFF", "Enable or disable the start-up-drift liveness check"],
        ["LOCK STATUS / LOCK LEARN <dev>", "Report lock state / start learn mode for an owner"],
        ["LOCK SET <dev> <ON,OFF,...>", "Store a pattern directly (4–12 presses)"],
        ["LOCK CLEAR / LOCK NOW / LOCK RELOCK <s>", "Remove pattern / lock immediately / auto-relock time"],
    ], "Main serial commands of the receiver firmware", widths=[7.0, 9.4], align="ll", size=10)

    R.table(["Path", "Content"], [
        ["firmware/esp32_receiver/src", "main.cpp (capture, protocol, outputs), phyauth_dsp.c (DSP + model), phyauth_lock.c (lock FSM), phyauth_model.h (exported model)"],
        ["firmware/esp32_transmitter/src", "main.cpp (RMT waveform, profiles, clone/emulation commands)"],
        ["dsp/pulse_dsp.py", "Clock recovery, decoding, quality gate, features (mirrored in C)"],
        ["ml/", "data loading and augmentation, models, triplet loss, authenticator, liveness, training, evaluation, export"],
        ["acquisition/", "Serial link with watchdog, protocol normalisation, simulator"],
        ["backend/", "FastAPI hub and routers: hardware, recordings, models, devices, events, attack lab, lock, WebSocket"],
        ["frontend/", "Next.js dashboard: RF Fingerprinting and Smart Lock workspaces"],
        ["tests/", "35 tests: DSP, ML, liveness, C parity, lock FSM, API"],
    ], "Repository layout", widths=[5.2, 11.2], align="ll", size=9.5)

    R.save(OUT_DOCX)
    return {"toc": R.toc, "figs": R.figs, "tabs": R.tabs}


def Inches_(x):
    from docx.shared import Inches
    return Inches(x)


def to_pdf():
    script = HERE / "topdf.applescript"
    script.write_text('on run argv\n  set inPath to POSIX file (item 1 of argv)\n  set outPath to POSIX file (item 2 of argv)\n'
                      '  tell application "Pages"\n    set d to open inPath\n    delay 2\n    export d to outPath as PDF\n'
                      '    close d saving no\n  end tell\nend run\n')
    if OUT_PDF.exists():
        OUT_PDF.unlink()
    subprocess.run(["osascript", str(script), str(OUT_DOCX), str(OUT_PDF)], check=True)


def find_pages(lists):
    import fitz
    doc = fitz.open(str(OUT_PDF))
    texts = [pg.get_text().replace("­", "") for pg in doc]
    norm = lambda s: " ".join(s.split())
    ntexts = [norm(t) for t in texts]
    # body starts after the List of Abbreviations page
    start = next(i for i, t in enumerate(ntexts) if "List of Abbreviations" in t and "Abbreviation Meaning" in t)
    pages = {}
    for key in ("toc", "figs", "tabs"):
        for item in lists[key]:
            text = item[1] if isinstance(item, tuple) else item
            probe = norm(text)
            if key == "toc":
                if probe.startswith("Chapter "):
                    probe = probe.split(": ", 1)[1]
                probe = probe[:60]
            else:
                probe = probe.split(":")[0] + ":"
            front = {"Abstract": "Keywords", "List of Figures": "Figure 2.1:", "List of Tables": "Table 1.1:",
                     "List of Abbreviations": "Abbreviation Meaning"}
            lo = 1 if text in front else start
            for i in range(lo, len(ntexts)):
                if text in front and (front[text] not in ntexts[i] or (text != "Abstract" and "Contents" in ntexts[i]
                                                                       and text != "List of Abbreviations")):
                    continue
                if probe in ntexts[i]:
                    pages[text] = i + 1
                    break
    return pages, len(doc)


if __name__ == "__main__":
    empty = {"toc": [], "figs": [], "tabs": []}
    lists = build(empty, {})
    lists = build(lists, {})
    to_pdf()
    pages, n = find_pages(lists)
    final = build(lists, pages)
    to_pdf()
    pages2, n2 = find_pages(final)
    if pages2 != pages:
        build(final, pages2)
        to_pdf()
    missing = [t for _, t in final["toc"] if t not in pages2] + [t for t in final["figs"] + final["tabs"] if t not in pages2]
    print("pages:", n2, "| unresolved:", missing)
